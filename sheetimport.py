"""
sheetimport.py — read the client's own Excel BOQ, whatever its layout
=====================================================================
Built 29 September 2026 for **Import BOQ from Excel v1**, CLIENT_CHANGES.md §0,
thirty-fourth block.

**This module imports nothing from the app** — not `store`, not `pipeline`, not
`boq`. It is handed bytes and a filename and hands back plain dicts; the routes,
the staging collection and every word of HTML belong to `boqimport.py`. That is
what lets it be tested on workbooks built in memory, and what keeps a parser
from ever reaching into a record. `tests/test_import_directions.py` holds the
leaf. `openpyxl` and `xlrd` are imported INSIDE the two readers, never at module
level — `photo.py`'s arrangement for Pillow — so a server without them still
boots and `/boq/import` says the reader is unavailable.

What it guarantees, in the order a file meets it
------------------------------------------------
1. **The bytes decide the format, never the name.** `PK\\x03\\x04` is a zip and
   must be a macro-free Excel workbook; `D0 CF 11 E0` is an OLE2 compound file
   and must be an Excel 97-2003 workbook. Anything else — a CSV, a PDF, a
   renamed executable — is refused with a sentence. An `.xlsm`, and an `.xls`
   carrying a VBA project, are refused outright: nothing here runs a macro, but
   a macro-bearing file is not one this form takes.
2. **Size is bounded before anything is parsed.** 5 MB uploaded; for an .xlsx
   the zip's declared uncompressed sizes are summed first and the file is
   refused above 50 MB, so a zip bomb is refused by arithmetic rather than by
   running out of memory. (`zipfile` itself refuses a member that inflates past
   the size it declared.)
3. **XML goes through defusedxml.** openpyxl picks it up by itself when it is
   installed; an .xlsx is REFUSED while `openpyxl.xml.DEFUSEDXML` is False.
4. **Nothing is written to disk and the file is not kept.** It is read from a
   BytesIO and dropped. What survives is the staged GRID — cell values, trimmed
   to the used range — never the workbook.
5. **Formulas are read, never computed.** The value Excel saved is the value. A
   formula with no saved value is FLAGGED — never evaluated, never guessed.
6. **Nothing the user must decide is guessed.** A rate column that does not say
   which track it is on is left UNSET, so the user picks Supply or Installation.
   "9.3+1.5+6" is flagged, never added up. A word in a quantity ("NA", "I.R.")
   is left BLANK and flagged, never 0.
   ⚠ **Except "RO" (1 October 2026).** "RO", "R.O.", "R/O" and "Rate only" mean
   the sheet quotes a RATE and no quantity: such a line comes in at quantity 0
   with its rates kept, marked *rate only* — not a flag. A 0 quantity is what
   it means: nothing is approved to claim until a revision gives it one.

⚠ **`item_no` is TEXT and only ever a display label.** A number is rounded to
its format's fixed decimals (a `0.00` cell to 2), otherwise to at most 2, and
its trailing zeros are stripped (1 October 2026): `5.199999999999999` reads
"5.2", `4.0` reads "4", and a `0.00`-formatted 1.1 reads "1.1" where it used to
read "1.10". Duplicates stay duplicates, and nothing here parses one into a
key. `line_id` is minted by `boq._clean_lines()` on save, like any typed line.

⚠ **Only the TOTAL quantity is read (v1).** Floor / area split columns and
separate take-off tabs are ignored — a decision, recorded in CLIENT_CHANGES.md.

⚠ **6 October 2026 — the user chooses, this module ADVISES** (CLIENT_CHANGES.md
§0, forty-fourth block). `advise()` suggests a target per column with one
plain-English reason (§7 below), `detect_names()` reads the project and the
customer from the rows above the heading (§8), and `build()` reads a discount
% and a net rate (checked, never imported) and a Remark column. **Only the
TICKED columns** — those the mapping gives a target — are read into a line,
flagged, footed or quoted; whether a row is a total or a heading is still
decided from the whole row.
"""

import datetime
import hashlib
import io
import json
import math
import re
import zipfile

# =============================================================================
# LIMITS — tune here, not in a branch
# =============================================================================

# The upload itself. Same ceiling as an attachment and a profile photo.
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

# An .xlsx is a zip of XML. Its declared uncompressed size is summed BEFORE
# openpyxl sees it; the client's largest real workbook is 2.1 MB inflated.
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

# The staged grid, per sheet. A sheet past either is REFUSED, never truncated:
# a schedule cut at row 1,500 is a schedule missing its last lines with nothing
# on the page saying so.
MAX_GRID_ROWS = 1500
MAX_GRID_COLS = 40

# ⚠ **4,000, not the 500 the build brief named.** Twenty of the fifty-six
#   clauses in the client's own seeded schedule run past 500 characters (the
#   longest is 1,369), and a BOQ header line carries the whole specification
#   paragraph. A 500 cap would have cut the client's own clause text on import,
#   which INTRODUCTION.md §9 forbids. A cell longer than this is cut AND
#   flagged on its row — never silently.
MAX_CELL_CHARS = 4000

# How far a sheet is walked before it is refused as not a schedule. Counted in
# rows the reader yields, empty ones included, so a sheet formatted down to
# row 1,048,576 is refused in bounded time.
MAX_SCAN_ROWS = 100_000
MAX_SCAN_COLS = 2_000

# The whole staged workbook, in JSON characters. The selected sheet always
# fits (it is checked on its own); other sheets are staged while they fit and
# listed as not staged after that. `db.sync()` re-hashes every record on every
# request, so a staging row is paid for until it is deleted.
STAGE_BUDGET_CHARS = 2_000_000

# Where the header is looked for, and what counts as one.
HEADER_SCAN_ROWS = 40
MIN_HEADER_SCORE = 2
HEADER_LABEL_MAX = 60

# How many data rows the preview shows under the column dropdowns.
PREVIEW_ROWS = 20

# The totals check: the sheet's grand total against sum(qty x rate), per track.
TOTALS_TOLERANCE = 1.0

# A numeric item number in a cell with no FIXED number format is rounded to at
# most this many decimals before its trailing zeros are stripped (1 October
# 2026): the Iron Mountain sheet stores item 17.2 as 17.200000000000003.
# ⚠ An item written to three decimals in a General cell (1.125) is rounded too.
ITEM_DECIMALS = 2

UNAVAILABLE = ("The Excel reader is not installed on this server. Ask whoever "
               "looks after it to run: pip install -r requirements.txt")

# =============================================================================
# THE TARGETS a column can be mapped to
# =============================================================================
#
# ⚠ **Six more than the build brief listed, each for a stated reason:**
#   `install_base_rate` / `install_escalation_pct` — the client's Sify sheet
#   carries an installation base rate, and the brief's own reproduction test
#   compares rates; `supply_amount` / `install_amount` — a two-track sheet has
#   one amount column per track, and the totals check is per track; `amount`
#   stays as the brief's single check-only column (one track, or the combined
#   total of both). None of the amount columns is imported — the form computes
#   amounts from quantity and rate, always.
TARGETS = (
    ("",                       "— ignore —"),
    ("item_no",                "Item No."),
    ("description",            "Description"),
    ("qty",                    "Quantity (total)"),
    ("unit",                   "Unit"),
    ("supply_base_rate",       "Supply · base rate"),
    ("escalation_pct",         "Supply · escalation %"),
    ("supply_rate",            "Supply · unit rate"),
    # 6 October 2026 (CLIENT_CHANGES.md §0, forty-fourth block, R1/R3): a
    # discount % per track — a real field on the BOQ line — and the sheet's
    # own net rate, read only to CHECK unit × (1 − discount) against it.
    ("supply_disc_pct",        "Supply · discount %"),
    ("supply_net_rate",        "Supply · net rate (check only)"),
    ("supply_amount",          "Supply · amount (check only)"),
    ("install_base_rate",      "Installation · base rate"),
    ("install_escalation_pct", "Installation · escalation %"),
    ("install_rate",           "Installation · unit rate"),
    ("install_disc_pct",       "Installation · discount %"),
    ("install_net_rate",       "Installation · net rate (check only)"),
    ("install_amount",         "Installation · amount (check only)"),
    ("amount",                 "Amount (check only)"),
    # 30 September 2026. The brand a line is priced on ("Jindal", "Newage")
    # has nowhere of its own on a BOQ line, so it goes into the line's
    # `remark` as "Make: Jindal" — captured, never printed, never lost.
    ("make",                   "Make (kept in the remark)"),
    # 6 October 2026 (R1). The sheet's own notes column, into the line's
    # `remark` — captured, never printed. Also where the advice sends a
    # discount written in rupees, which this pass does not read as a discount.
    ("remark",                 "Remark (internal, never printed)"),
)
TARGET_KEYS = tuple(k for k, _l in TARGETS)
TARGET_LABEL = dict(TARGETS)

# A guessed column the user MUST decide. Never a valid target on confirm.
UNDECIDED = "?"

RATE_FIELDS = ("supply_base_rate", "escalation_pct", "supply_rate",
               "install_base_rate", "install_escalation_pct", "install_rate")
AMOUNT_FIELDS = ("supply_amount", "install_amount", "amount")
DISC_FIELDS = ("supply_disc_pct", "install_disc_pct")
NET_FIELDS = ("supply_net_rate", "install_net_rate")
PCT_FIELDS = ("escalation_pct", "install_escalation_pct") + DISC_FIELDS


def available() -> dict:
    """{"xlsx": bool, "xls": bool} — whether each reader can run on this
    server. An .xlsx reader without defusedxml counts as unavailable, because
    `_read_xlsx()` refuses in that state."""
    out = {}
    try:
        import openpyxl.xml
        out["xlsx"] = bool(getattr(openpyxl.xml, "DEFUSEDXML", False))
    except ImportError:
        out["xlsx"] = False
    try:
        import xlrd  # noqa: F401
        out["xls"] = True
    except ImportError:
        out["xls"] = False
    return out


class Refused(Exception):
    """A file (or a sheet) this module will not read, with the sentence the
    operator is shown. Never carries a stack trace or a library's wording."""

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


# =============================================================================
# SMALL HELPERS
# =============================================================================

def col_letter(i: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA. The column name the operator sees in Excel."""
    s, n = "", int(i) + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _norm(label) -> str:
    """A header cell, lower-cased and whitespace-collapsed."""
    return re.sub(r"\s+", " ", str(label or "")).strip().lower()


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _kind(grid: dict, ri: int, ci) -> str:
    if ci is None:
        return ""
    return (grid.get("kinds") or {}).get(f"{ri}:{ci}", "")


def _value(grid: dict, ri: int, ci):
    if ci is None:
        return None
    vals = grid["rows"][ri][1]
    return vals[ci] if ci < len(vals) else None


# =============================================================================
# 1. THE BYTES — format, size, macros
# =============================================================================

XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
# How the VBA storage is named in an OLE2 directory entry (UTF-16LE).
_VBA_OLE_NAME = "_VBA_PROJECT_CUR".encode("utf-16-le")

MACRO_REFUSAL = ("This workbook carries macros (.xlsm). Open it in Excel, "
                 "use Save As → Excel Workbook (.xlsx), and upload that copy.")


def sniff_format(data: bytes) -> str:
    """
    `"xlsx"` or `"xls"`, from the LEADING BYTES — or `Refused`.

    The filename and the browser's Content-Type are the uploader's to choose;
    the content is not. `attachment.sniff()`'s rule.
    """
    if data.startswith(XLSX_MAGIC):
        return "xlsx"
    if data.startswith(XLS_MAGIC):
        return "xls"
    raise Refused("That file is not an Excel workbook. Upload the BOQ as an "
                  ".xlsx or .xls file — a CSV, PDF or scanned sheet cannot be "
                  "read here.")


def _check_xlsx_container(data: bytes) -> None:
    """The zip-level refusals, all decided before openpyxl is called."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise Refused("That file looks like an .xlsx but it could not be "
                      "opened — it may be damaged. Open it in Excel, save it "
                      "again and upload the new copy.")
    total = sum(max(0, int(i.file_size)) for i in infos)
    if total > MAX_UNCOMPRESSED_BYTES:
        raise Refused(f"That workbook unpacks to {total / 1048576:.0f} MB, and "
                      f"the limit is {MAX_UNCOMPRESSED_BYTES // 1048576} MB. "
                      f"Save the BOQ sheet on its own as a new workbook and "
                      f"upload that.")
    names = [i.filename.lower() for i in infos]
    if any(n.endswith("vbaproject.bin") for n in names):
        raise Refused(MACRO_REFUSAL)
    try:
        ctypes_xml = zf.read("[Content_Types].xml").decode("utf-8", "replace").lower()
    except KeyError:
        raise Refused("That file is a zip archive but not an Excel workbook.")
    if "macroenabled" in ctypes_xml:
        raise Refused(MACRO_REFUSAL)
    if ("spreadsheetml.sheet.main+xml" not in ctypes_xml
            and "spreadsheetml.template.main+xml" not in ctypes_xml):
        raise Refused("That file is an Office document but not an Excel "
                      "workbook. Upload the BOQ as an .xlsx or .xls file.")


# =============================================================================
# 2. READING — each reader yields the same neutral shape
# =============================================================================
#
# A sheet is read into {row_number: {col_index: (value, kind)}} where value is
# a str, an int/float, or "" for a special cell, and kind is:
#
#   ""        an ordinary value
#   "error"   an Excel error (#VALUE!, #REF!, …) — value is its text
#   "date"    a date-formatted cell — value is its ISO text
#   "formula" a formula with NO saved value — value is ""
#   "pct"     a number formatted as a percentage (0.15 shown as 15%)
#   "dp:N"    a number formatted with N fixed decimals (item "1.10")
#   "bool"    TRUE / FALSE
#   "long"    text cut at MAX_CELL_CHARS

def _num_kind(number_format: str) -> str:
    fmt = str(number_format or "")
    if "%" in fmt:
        return "pct"
    m = re.fullmatch(r"0\.(0+)", fmt.strip())
    return f"dp:{len(m.group(1))}" if m else ""


def _text_cell(v: str):
    s = v if isinstance(v, str) else str(v)
    if not s.strip():
        return None
    if len(s) > MAX_CELL_CHARS:
        return s[:MAX_CELL_CHARS], "long"
    return s, ""


class _SheetTooBig(Exception):
    pass


class _Collector:
    """Accumulates one sheet's non-empty cells, refusing past the limits."""

    def __init__(self):
        self.cells = {}
        self.cols = set()

    def put(self, r: int, c: int, value, kind: str) -> None:
        row = self.cells.get(r)
        if row is None:
            if len(self.cells) >= MAX_GRID_ROWS:
                raise _SheetTooBig(
                    f"more than {MAX_GRID_ROWS:,} rows with something in them")
            row = self.cells[r] = {}
        row[c] = (value, kind)
        if c not in self.cols:
            self.cols.add(c)
            if len(self.cols) > MAX_GRID_COLS:
                raise _SheetTooBig(
                    f"more than {MAX_GRID_COLS} columns with something in them")


def _read_xlsx(data: bytes) -> list:
    """[(name, visibility, cells_or_refusal)] for every worksheet."""
    _check_xlsx_container(data)
    try:
        import openpyxl
        import openpyxl.xml
    except ImportError:
        raise Refused(UNAVAILABLE)
    if not getattr(openpyxl.xml, "DEFUSEDXML", False):
        # Fail closed. openpyxl would otherwise parse the workbook's XML with
        # the standard library, and this is a file somebody else wrote.
        raise Refused("The Excel reader on this server is missing its XML "
                      "safety layer (defusedxml), so .xlsx files are refused. "
                      "Ask whoever looks after the server to run: "
                      "pip install -r requirements.txt")

    try:
        wb_v = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                      data_only=True, keep_links=False)
        wb_f = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                      data_only=False, keep_links=False)
    except Exception as exc:                       # openpyxl raises many shapes
        raise Refused(_unreadable(exc, "an .xlsx"))

    out = []
    try:
        for ws_v, ws_f in zip(wb_v.worksheets, wb_f.worksheets):
            name = str(ws_v.title)
            vis = str(getattr(ws_v, "sheet_state", "visible") or "visible")
            try:
                out.append((name, vis, _xlsx_cells(ws_v, ws_f)))
            except _SheetTooBig as exc:
                out.append((name, vis, str(exc)))
    except Exception as exc:
        raise Refused(_unreadable(exc, "an .xlsx"))
    finally:
        for wb in (wb_v, wb_f):
            try:
                wb.close()
            except Exception:
                pass
    return out


def _xlsx_cells(ws_v, ws_f) -> dict:
    """One worksheet, read from the two parses in lockstep: values from the
    data-only parse, and "is this a formula" from the other. openpyxl cannot
    give both from one parse."""
    ws_v.reset_dimensions()
    ws_f.reset_dimensions()
    col = _Collector()
    f_rows = ws_f.iter_rows()
    for r, row_v in enumerate(ws_v.iter_rows(), start=1):
        row_f = next(f_rows, ())
        if r > MAX_SCAN_ROWS:
            raise _SheetTooBig(f"it runs past row {MAX_SCAN_ROWS:,}")
        if len(row_v) > MAX_SCAN_COLS:
            raise _SheetTooBig(f"a row runs past column {MAX_SCAN_COLS:,}")
        for c, cell in enumerate(row_v):
            v = cell.value
            if v is None:
                fc = row_f[c] if c < len(row_f) else None
                if fc is not None and getattr(fc, "data_type", None) == "f":
                    col.put(r, c, "", "formula")
                continue
            dt = getattr(cell, "data_type", None)
            if dt == "e":
                col.put(r, c, str(v), "error")
            elif isinstance(v, bool):
                col.put(r, c, "TRUE" if v else "FALSE", "bool")
            elif isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
                col.put(r, c, v.isoformat(), "date")
            elif isinstance(v, (int, float)):
                if not math.isfinite(v):
                    col.put(r, c, "#NUM!", "error")
                elif getattr(cell, "is_date", False):
                    col.put(r, c, str(v), "date")
                else:
                    col.put(r, c, v, _num_kind(getattr(cell, "number_format", "")))
            else:
                t = _text_cell(v)
                if t is not None:
                    col.put(r, c, t[0], t[1])
    return col


def _read_xls(data: bytes) -> list:
    """[(name, visibility, cells_or_refusal)] for every sheet of an .xls."""
    if _VBA_OLE_NAME in data:
        raise Refused(MACRO_REFUSAL.replace("(.xlsm)", "(a VBA project)"))
    try:
        import xlrd
    except ImportError:
        raise Refused(UNAVAILABLE)
    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True,
                                  on_demand=True)
    except Exception as exc:                       # xlrd raises many shapes
        raise Refused(_unreadable(exc, "an Excel 97-2003 (.xls)"))

    vis_name = {0: "visible", 1: "hidden", 2: "veryHidden"}
    out = []
    try:
        for i in range(book.nsheets):
            sh = book.sheet_by_index(i)
            vis = vis_name.get(getattr(sh, "visibility", 0), "visible")
            try:
                out.append((str(sh.name), vis, _xls_cells(book, sh, xlrd)))
            except _SheetTooBig as exc:
                out.append((str(sh.name), vis, str(exc)))
            book.unload_sheet(i)
    except Exception as exc:
        raise Refused(_unreadable(exc, "an Excel 97-2003 (.xls)"))
    finally:
        try:
            book.release_resources()
        except Exception:
            pass
    return out


def _xls_cells(book, sh, xlrd):
    if sh.nrows > MAX_SCAN_ROWS:
        raise _SheetTooBig(f"it runs past row {MAX_SCAN_ROWS:,}")
    col = _Collector()
    for r in range(sh.nrows):
        n = sh.row_len(r)
        if n > MAX_SCAN_COLS:
            raise _SheetTooBig(f"a row runs past column {MAX_SCAN_COLS:,}")
        for c in range(n):
            t = sh.cell_type(r, c)
            if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                continue
            v = sh.cell_value(r, c)
            if t == xlrd.XL_CELL_TEXT:
                tc = _text_cell(v)
                if tc is not None:
                    col.put(r + 1, c, tc[0], tc[1])
            elif t == xlrd.XL_CELL_NUMBER:
                if not math.isfinite(v):
                    col.put(r + 1, c, "#NUM!", "error")
                else:
                    col.put(r + 1, c, v, _num_kind(_xls_format(book, sh, r, c)))
            elif t == xlrd.XL_CELL_DATE:
                try:
                    txt = xlrd.xldate_as_datetime(v, book.datemode).isoformat()
                except Exception:
                    txt = str(v)
                col.put(r + 1, c, txt, "date")
            elif t == xlrd.XL_CELL_BOOLEAN:
                col.put(r + 1, c, "TRUE" if v else "FALSE", "bool")
            elif t == xlrd.XL_CELL_ERROR:
                col.put(r + 1, c, xlrd.error_text_from_code.get(v, "#ERROR"), "error")
    return col


def _xls_format(book, sh, r, c) -> str:
    try:
        xf = book.xf_list[sh.cell_xf_index(r, c)]
        return book.format_map[xf.format_key].format_str
    except Exception:
        return ""


def _unreadable(exc, what: str) -> str:
    """The refusal for a file a reader choked on — never the library's words."""
    if type(exc).__module__.startswith("defusedxml"):
        return ("That workbook contains XML this reader refuses to process "
                "(entity or DTD declarations). Open it in Excel, save it as a "
                "new .xlsx and upload that copy.")
    return (f"That file could not be read as {what} workbook — it may be "
            f"damaged or password-protected. Open it in Excel, save it again "
            f"and upload the new copy.")


# =============================================================================
# 3. STAGING — the grid, trimmed to the used range
# =============================================================================

def _grid_of(col: _Collector) -> dict:
    """
    The collector as a staged grid.

    `cols` is the ORIGINAL 0-based column index of each kept column, so a
    column is always named by the letter the operator sees in Excel; `rows`
    carries the ORIGINAL row number for the same reason — a flag that says
    "row 57" has to mean row 57 of their sheet. Only rows and columns with
    something in them are kept.
    """
    cols = sorted(col.cols)
    pos = {c: i for i, c in enumerate(cols)}
    rows, kinds = [], {}
    for r in sorted(col.cells):
        vals = [None] * len(cols)
        for c, (v, k) in col.cells[r].items():
            ci = pos[c]
            vals[ci] = v
            if k:
                kinds[f"{len(rows)}:{ci}"] = k
        while vals and vals[-1] is None:
            vals.pop()
        rows.append([r, vals])
    grid = {"cols": cols, "rows": rows, "kinds": kinds}
    grid["header"] = detect_header(grid)
    return grid


def read(data: bytes, filename: str = "") -> dict:
    """
    Read an uploaded workbook into the staged shape, or raise `Refused`.

    Returns::

        {"format": "xlsx" | "xls",
         "sheets": [{"name", "visibility", "rows", "cols", "staged",
                     "refusal", "score"}, …],     # in workbook order
         "grid":   [grid | None, …],               # aligned with "sheets"
         "selected": index of the pre-selected sheet}

    ⚠ The pre-selected sheet is the VISIBLE sheet with the most rows carrying a
      number in both a quantity-like and a rate-like column. A hidden sheet is
      listed and may be chosen, and is never chosen for the operator.
    """
    if not data:
        raise Refused("That file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise Refused(f"That file is {len(data) / 1048576:.1f} MB and the "
                      f"limit is {MAX_UPLOAD_BYTES // 1048576} MB. Save the BOQ "
                      f"sheet on its own as a new workbook and upload that.")
    fmt = sniff_format(data)
    if fmt == "xlsx" and str(filename or "").lower().endswith(".xlsm"):
        raise Refused(MACRO_REFUSAL)
    raw = _read_xlsx(data) if fmt == "xlsx" else _read_xls(data)
    return _assemble(fmt, raw)


def from_rows(sheets: list, fmt: str = "xlsx") -> dict:
    """
    `read()`'s result from plain rows — `[(name, visibility, [[cell, …], …])]`
    — with no workbook and no reader involved.

    It exists so a caller with no openpyxl (ABOUT.md §1's row-3 configuration)
    can still produce a staged workbook through the SAME staging, header,
    scoring and budget code `read()` uses: `tests/test_entity_fallbacks.py`'s
    sweep fixture is that caller. Cells are taken as plain values; there are
    no formats, dates or formulas to carry.
    """
    raw = []
    for name, vis, rows in sheets:
        col = _Collector()
        try:
            for r, vals in enumerate(rows, start=1):
                for c, v in enumerate(vals):
                    if v is None:
                        continue
                    if _is_number(v):
                        col.put(r, c, v, "")
                    else:
                        t = _text_cell(v)
                        if t is not None:
                            col.put(r, c, t[0], t[1])
        except _SheetTooBig as exc:
            raw.append((str(name), str(vis), str(exc)))
            continue
        raw.append((str(name), str(vis), col))
    return _assemble(fmt, raw)


def _assemble(fmt: str, raw: list) -> dict:
    """The per-sheet staging, scoring, pre-selection and budget that turn a
    reader's output into what `read()` returns."""
    sheets, grids = [], []
    for name, vis, cells in raw:
        entry = {"name": name, "visibility": vis, "rows": 0, "cols": 0,
                 "staged": False, "refusal": "", "score": 0, "qty_rows": 0}
        grid = None
        if isinstance(cells, str):
            entry["refusal"] = f"Not read: {cells}. Split the sheet and import each part."
        else:
            grid = _grid_of(cells)
            entry["rows"], entry["cols"] = len(grid["rows"]), len(grid["cols"])
            entry["score"], entry["qty_rows"] = _schedule_score(grid)
            entry["staged"] = bool(grid["rows"])
            if not grid["rows"]:
                entry["refusal"] = "Empty."
                grid = None
        sheets.append(entry)
        grids.append(grid)

    if not any(s["staged"] for s in sheets):
        why = "; ".join(f"{s['name']}: {s['refusal']}" for s in sheets) or "no sheets"
        raise Refused(f"No sheet in that workbook could be read ({why}).")

    selected = pick_sheet(sheets)

    # The budget. The selected sheet is checked on its own; the rest ride
    # while they fit, in workbook order, and are listed as not staged after.
    used = 0
    for i in [selected] + [j for j in range(len(sheets)) if j != selected]:
        if grids[i] is None:
            continue
        size = len(json.dumps(grids[i], ensure_ascii=False, separators=(",", ":")))
        if used + size > STAGE_BUDGET_CHARS:
            if i == selected:
                raise Refused(f"The sheet “{sheets[i]['name']}” holds more text "
                              f"than one import can carry. Split it into two "
                              f"sheets and import each.")
            sheets[i]["staged"] = False
            sheets[i]["refusal"] = ("Not held for preview — the workbook is too "
                                    "large to stage every sheet. Save this sheet "
                                    "on its own and import it.")
            grids[i] = None
            continue
        used += size
    return {"format": fmt, "sheets": sheets, "grid": grids, "selected": selected}


def pick_sheet(sheets: list) -> int:
    """
    The visible staged sheet with the highest schedule score; first wins a tie.

    ⚠ The score is the brief's — rows with a number in BOTH a quantity-like
      and a rate-like column — and `qty_rows` (rows with a quantity at all) only
      breaks a tie. Found on the client's own files: a workbook whose rate
      cells are empty scores 0 on every sheet, and the first sheet — its
      one-page SUMMARY — was chosen over the 131-row schedule behind it.

    Falls back to the first staged sheet, visible or not, only when no visible
    sheet could be staged at all.
    """
    def rank(s):
        return (s.get("score", 0), s.get("qty_rows", 0))

    best = None
    for i, s in enumerate(sheets):
        if s.get("staged") and s.get("visibility") == "visible":
            if best is None or rank(s) > rank(sheets[best]):
                best = i
    if best is None:
        best = next(i for i, s in enumerate(sheets) if s.get("staged"))
    return best


# =============================================================================
# 4. THE HEADER — where it is, and what each column probably is
# =============================================================================
#
# The keywords the brief named, plus four this pass added and records:
# "sl" (Sl. No. is on six of the client's own sheets), "specification",
# "material" / "labour" (the common names for the two tracks in Indian
# schedules). Nothing here reads a header longer than 60 characters: a long
# cell is data, not a column head.

_KW = {
    "sr":           re.compile(r"\bsr\b|\bs\.?\s?no\b|\bsl\b|\bserial\b"),
    "item":         re.compile(r"\bitems?\b"),
    "description":  re.compile(r"descr|particular|specification|name of (?:the )?(?:item|work)"),
    "qty":          re.compile(r"\bqty\b|quantit|\bqnty\b"),
    "unit":         re.compile(r"\bunits?\b|\buom\b|u\.o\.m"),
    "rate":         re.compile(r"\brates?\b"),
    "amount":       re.compile(r"\bamount\b|\bamt\b"),
    "supply":       re.compile(r"\bsupply\b|\bmaterial\b"),
    "installation": re.compile(r"install|\blabou?r\b|\berection\b"),
    "base":         re.compile(r"\bbase\b|\bbasic\b"),
    "esc":          re.compile(r"\besc\b|escal"),
}
_TOTAL_WORD = re.compile(r"\btotal\b")


def _hits(label: str) -> set:
    t = _norm(label)
    if not t or len(t) > HEADER_LABEL_MAX:
        return set()
    return {k for k, rx in _KW.items() if rx.search(t)}


def _labels(vals: list) -> list:
    out = []
    for v in vals:
        out.append(v if isinstance(v, str) else "")
    return out


def _score(labels: list) -> int:
    return sum(len(_hits(l)) for l in labels)


def _combine(upper: list, lower: list):
    """
    A two-row header, merged: "Supply" over "Rate | Amount" becomes
    "supply rate | supply amount".

    An upper label carries rightward across empty upper cells for as long as
    the lower row has a label — the shape a merged cell leaves in the grid.
    ⚠ Only when the upper row has TWO or more labels: a one-label upper row is
      a title ("BILL OF QUANTITIES"), and carrying it across every column would
      stamp "quantities" on all of them.
    ⚠ And only when the lower row is MADE OF heading words: no figure in it,
      and every label either a keyword, the word "total", or three characters
      at most (a floor code: "GF", "B1"). Found by the suite: the first DATA
      row, carrying "Rate Only" in its quantity, scored a keyword and was
      swallowed into the header — the line vanished with nothing said.
    """
    # ⚠ **5 October 2026 (CLIENT_CHANGES.md §0 forty-third block, R2a).** A
    #   lower row that carries a RATE or AMOUNT label under a heading ("Unit
    #   Rate" under "Installation") is a header row even when it also carries
    #   floor or area labels longer than three characters ("Terrace", "Riser"
    #   under "DC building") — the client's Nxtra sheet, whose rate column was
    #   never mapped because "Terrace" failed the rule below. Such a row still
    #   may hold NO figure, and each of its other labels is held to the header
    #   label length, so a data row cannot pass for one.
    relaxed = _rate_label_under_heading(upper, lower)
    for v in lower:
        if _is_number(v):
            return None
        if isinstance(v, str) and v.strip():
            t = _norm(v)
            if relaxed and len(t) <= HEADER_LABEL_MAX:
                continue
            if not (_hits(t) or _TOTAL_WORD.search(t) or len(t) <= 3):
                return None
    up, lo = _labels(upper), _labels(lower)
    width = max(len(up), len(lo))
    up += [""] * (width - len(up))
    lo += [""] * (width - len(lo))
    if sum(1 for u in up if u.strip() and len(u) <= HEADER_LABEL_MAX) < 2:
        return None
    out, carry = [], ""
    for u, l in zip(up, lo):
        if u.strip():
            carry = u
        elif not l.strip():
            carry = ""
        out.append(f"{carry} {l}".strip() if u.strip() or l.strip() else "")
    return out


# A cell that is nothing but a rate or an amount heading — "Unit Rate", "Rate",
# "Amount", "Total cost", with or without "(INR)" (R2a, 5 October 2026).
_RATE_AMOUNT_LABEL = re.compile(
    r"^\s*(?:unit\s*rate|rate|amount|total\s*cost)"
    r"(?:\s*\(?\s*(?:inr|rs\.?|₹)\s*\)?)?\s*$", re.I)


def _rate_label_under_heading(upper: list, lower: list) -> bool:
    """
    Does the lower row carry a rate or amount label UNDER a heading — a
    non-empty upper cell in its column, or one carried rightward to it the way
    `_combine()` carries a merged cell? (R2a.)
    """
    up, lo = _labels(upper), _labels(lower)
    width = max(len(up), len(lo))
    up += [""] * (width - len(up))
    lo += [""] * (width - len(lo))
    carry = ""
    for u, l in zip(up, lo):
        if u.strip():
            carry = u
        elif not l.strip():
            carry = ""
        if carry.strip() and _RATE_AMOUNT_LABEL.match(l or ""):
            return True
    return False


# The tokens a units row under a header carries and nothing else — "INR | INR"
# under "Unit Rate | Total cost" (R2b, 5 October 2026). Compared lower-cased
# and stripped.
_UNIT_TOKENS = {"inr", "rs", "rs.", "₹", "(inr)", "nos", "%"}


def _units_only(vals: list) -> bool:
    """A row whose non-empty cells are ONLY currency or unit tokens, and at
    least one of them. Such a row directly under the header belongs to it."""
    seen = False
    for v in vals:
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if not isinstance(v, str) or v.strip().lower() not in _UNIT_TOKENS:
            return False
        seen = True
    return seen


def _heading_like(vals: list) -> bool:
    """
    Does this row, on its own, read as a column heading?

    At least MIN_HEADER_SCORE keyword hits spread over at least TWO cells. A
    single cell is a title ("SUPPLY & INSTALLATION RATES" scores three hits in
    one cell and is not a heading), which is why the spread is required.
    """
    labels = _labels(vals)
    hit_cells = sum(1 for l in labels if _hits(l))
    return hit_cells >= 2 and _score(labels) >= MIN_HEADER_SCORE


def detect_header(grid: dict) -> list:
    """
    `[top, bottom]` grid-row indexes of the header (equal for a one-row
    header), or `[]` when no row in the first 40 reads as one.

    ⚠ **The FIRST heading row wins (30 September 2026), not the best-scoring
      one.** The client's Sify sheet prints its column heading at row 2 and
      again at row 28, above section B; the repeat scores 11 against the
      original's 9, so "most hits wins" took row 28 and every line above it —
      the whole of section A, 18 lines — was dropped with nothing said. A
      later repeat is now skipped as a repeated heading by `build()`.

    A pair ("Supply" over "Rate | Amount") is taken when the two rows combined
    beat each alone, so a header with a section row under it stays one row.

    ⚠ **A units row directly under the band joins it (5 October 2026, R2b)**
      — a row whose cells are ONLY currency or unit tokens ("INR | INR" under
      "Unit Rate | Total cost"). It was read as a priced line with no
      description and no quantity. The band then ends on that row, so
      `data_start()` skips it; `header_labels()` reads the label rows only.
    """
    rows = grid["rows"]

    def with_units(band: list) -> list:
        nxt = band[1] + 1
        if (nxt < len(rows) and rows[nxt][0] == rows[band[1]][0] + 1
                and _units_only(rows[nxt][1])):
            return [band[0], nxt]
        return band

    for ri, (rnum, vals) in enumerate(rows):
        if rnum > HEADER_SCAN_ROWS:
            break
        pair = None
        if ri + 1 < len(rows) and rows[ri + 1][0] == rnum + 1:
            combined = _combine(vals, rows[ri + 1][1])
            if combined is not None:
                s2 = _score(combined)
                if s2 > max(_score(_labels(vals)), _score(_labels(rows[ri + 1][1]))) \
                        and s2 >= MIN_HEADER_SCORE:
                    pair = [ri, ri + 1]
        if _heading_like(vals):
            return with_units(pair or [ri, ri])
        if pair and _heading_like(combined):
            return with_units(pair)
    return []


def _label_band(grid: dict) -> list:
    """The header band without a trailing units row — the rows that carry
    the column LABELS (R2b). `[]` when there is no header."""
    hdr = list(grid.get("header") or [])
    if hdr and hdr[1] > hdr[0] and _units_only(grid["rows"][hdr[1]][1]):
        hdr[1] -= 1
    return hdr


def header_labels(grid: dict) -> dict:
    """{grid column index: the column's header text} — combined for a pair."""
    hdr = _label_band(grid)
    if not hdr:
        return {}
    top = grid["rows"][hdr[0]][1]
    if hdr[1] != hdr[0]:
        labels = _combine(top, grid["rows"][hdr[1]][1]) or _labels(grid["rows"][hdr[1]][1])
    else:
        labels = _labels(top)
    return {ci: l for ci, l in enumerate(labels) if l.strip()}


def signature(grid: dict) -> str:
    """
    The LAYOUT SIGNATURE — sha256 of the normalised header texts and the
    column each sits in. Two uploads of the same template produce the same
    signature whatever their data, which is what lets a confirmed mapping be
    reused. "" when no header was found: a sheet with no header is never a
    known layout.
    """
    labels = header_labels(grid)
    if not labels:
        return ""
    cols = grid["cols"]
    payload = json.dumps(sorted((int(cols[ci]), _norm(l)) for ci, l in labels.items()),
                         ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def guess_mapping(grid: dict) -> dict:
    """
    {str(Excel column index): target | "?" | ""} — a first guess the preview
    shows under every column, which the user confirms or changes.

    ⚠ **A rate column is given a TRACK only when its own header, or the band
      above it, names the track** ("Supply Rate", "Installation" over "Rate").
      A rate column that does not is left "?" — the user must choose Supply or
      Installation, and confirm refuses until they do. With one rate on the
      sheet that is the brief's rule exactly; with several it is the same rule
      applied to each.
    """
    labels = header_labels(grid)
    cols = grid["cols"]
    info = {}
    for ci, label in labels.items():
        h = _hits(label)
        track = ("supply" if "supply" in h else "install" if "installation" in h else "")
        info[ci] = (h, track, _norm(label))

    out = {str(c): "" for c in cols}
    taken = set()

    def put(ci, target):
        out[str(cols[ci])] = target
        if target and target != UNDECIDED:
            taken.add(target)

    order = sorted(info)
    # Item number: a serial column first; "Item" alone only when there is none
    # and it is not the description ("Item Description").
    for ci in order:
        h, _t, _l = info[ci]
        if "sr" in h and "description" not in h and "item_no" not in taken:
            put(ci, "item_no")
    for ci in order:
        h, _t, _l = info[ci]
        if (out[str(cols[ci])] == "" and "description" in h
                and "description" not in taken):
            put(ci, "description")
    if "item_no" not in taken:
        for ci in order:
            h, _t, _l = info[ci]
            if out[str(cols[ci])] == "" and "item" in h and not (h & {"rate", "amount", "qty", "unit"}):
                put(ci, "item_no")
                break
    # Quantity: the one that says "total" when there are several (floor /
    # area splits are ignored in v1); otherwise the first.
    qtys = [ci for ci in order if out[str(cols[ci])] == ""
            and "qty" in info[ci][0] and not (info[ci][0] & {"rate", "amount"})]
    if qtys:
        tot = [ci for ci in qtys if _TOTAL_WORD.search(info[ci][2])]
        put((tot or qtys)[0], "qty")
    for ci in order:
        h, _t, _l = info[ci]
        if (out[str(cols[ci])] == "" and "unit" in h
                and not (h & {"rate", "amount", "qty"}) and "unit" not in taken):
            put(ci, "unit")

    # Rates and escalations, per track.
    rates, escs = {}, {}
    for ci in order:
        h, track, _l = info[ci]
        if out[str(cols[ci])] != "":
            continue
        if "esc" in h:
            escs.setdefault(track, []).append(ci)
        elif "rate" in h and "amount" not in h:
            rates.setdefault(track, []).append(ci)
    for track in sorted(set(rates) | set(escs)):
        rs, es = rates.get(track, []), escs.get(track, [])
        if not track:
            # No track named anywhere near these columns: every one is the
            # user's decision.
            for ci in rs + es:
                put(ci, UNDECIDED)
            continue
        base_t = "supply_base_rate" if track == "supply" else "install_base_rate"
        esc_t = "escalation_pct" if track == "supply" else "install_escalation_pct"
        unit_t = "supply_rate" if track == "supply" else "install_rate"
        if es:
            put(es[0], esc_t)
            left = [ci for ci in rs if ci < es[0]]
            right = [ci for ci in rs if ci > es[0]]
            if left:
                put(left[-1], base_t)
            if right:
                put(right[0], unit_t)
        elif len(rs) >= 2 and any("base" in info[ci][0] for ci in rs):
            b = next(ci for ci in rs if "base" in info[ci][0])
            put(b, base_t)
            put(next(ci for ci in rs if ci != b), unit_t)
        elif rs:
            put(rs[0], unit_t)

    # Amounts — check only. One per track, and one without a track.
    #
    # ⚠ An Amount column whose heading names no track takes the track of the
    #   RATE column directly to its left (30 September 2026) — "Supply | Amount |
    #   Installation | Amount" is the Jamnagar sheet's shape, where the second
    #   Amount used to be ignored and the first read as a combined figure. Only
    #   when the heading itself is silent, and only from a rate column already
    #   given a track; the preview's dropdown shows the result and can undo it.
    rate_track = {"supply_rate": "supply", "install_rate": "install"}
    for ci in order:
        h, track, label = info[ci]
        if out[str(cols[ci])] != "":
            continue
        is_amount = "amount" in h or (_TOTAL_WORD.search(label) and "qty" not in h)
        if not is_amount:
            continue
        if not track and ci > 0:
            track = rate_track.get(out[str(cols[ci - 1])], "")
        t = ("supply_amount" if track == "supply"
             else "install_amount" if track == "install" else "amount")
        if t not in taken:
            put(ci, t)

    # Make / brand — kept in the line's remark.
    for ci in order:
        if out[str(cols[ci])] == "" and "make" not in taken and _MAKE.search(info[ci][2]):
            put(ci, "make")
    return out


_MAKE = re.compile(r"^(?:make|brand|makes|manufacturer|mfr\.?|make\s*/\s*brand)$")


def _schedule_score(grid: dict) -> tuple:
    """`(score, qty_rows)` — rows under the header with a number in the guessed
    quantity column AND in some rate-like column, and rows with a quantity at
    all. What `pick_sheet()` ranks by, in that order."""
    if not grid.get("header"):
        return 0, 0
    mapping = guess_mapping(grid)
    cols = grid["cols"]
    pos = {str(c): i for i, c in enumerate(cols)}
    qty = [pos[k] for k, t in mapping.items() if t == "qty"]
    rate_like = [pos[k] for k, t in mapping.items()
                 if t in RATE_FIELDS or t == UNDECIDED]
    if not qty:
        return 0, 0
    n = q = 0
    for ri in range(grid["header"][1] + 1, len(grid["rows"])):
        if not (_is_number(_value(grid, ri, qty[0])) and _kind(grid, ri, qty[0]) != "date"):
            continue
        q += 1
        if any(_is_number(_value(grid, ri, c)) for c in rate_like):
            n += 1
    return n, q


def data_start(grid: dict) -> int:
    """The first grid row under the header (0 when there is no header)."""
    hdr = grid.get("header") or []
    return hdr[1] + 1 if hdr else 0


# The words in a heading that say the sheet's rates are what WE pay, not what
# we sell at (1 October 2026 — the Iron Mountain sheet is headed "OWN COST
# (INR)"). The brief's four, plus their plural / -ing forms. "own cost" is
# tried first so it is reported as itself and not as a bare "cost".
_COST_WORDS = re.compile(r"\b(own\s+costs?|costs?|buy(?:ing)?|purchases?)\b", re.I)


def cost_words(grid: dict) -> list:
    """
    `[(word, cell text), …]` — every cost word in the HEADING ROWS, first
    occurrence of each, in the order the heading reads. `[]` when the sheet has
    no recognised heading or the heading says nothing about cost.

    It only PRE-SELECTS "Our cost" on the preview, with these words shown as
    the reason; the operator can change it either way. A heading that says
    "cost" and means a selling price is exactly why it is never decided here.
    """
    hdr = grid.get("header") or []
    if not hdr:
        return []
    out, seen = [], set()
    for ri in range(hdr[0], hdr[1] + 1):
        for v in grid["rows"][ri][1]:
            if not isinstance(v, str):
                continue
            for m in _COST_WORDS.finditer(v):
                w = re.sub(r"\s+", " ", m.group(1)).lower()
                if w not in seen:
                    seen.add(w)
                    out.append((w, re.sub(r"\s+", " ", v).strip()))
    return out


# =============================================================================
# 5. THE MAPPING — what confirm refuses
# =============================================================================

def clean_mapping(grid: dict, posted: dict) -> dict:
    """A posted {column: target} reduced to known columns and known targets.
    An unknown target reads as ignore; "?" is kept so it can be refused."""
    out = {}
    for c in grid["cols"]:
        t = str(posted.get(str(c), "") or "")
        out[str(c)] = t if (t in TARGET_KEYS or t == UNDECIDED) else ""
    return out


def mapping_problems(grid: dict, mapping: dict) -> list:
    """Why this mapping cannot be confirmed yet — [] when it can."""
    probs = []
    for c in grid["cols"]:
        if mapping.get(str(c)) == UNDECIDED:
            probs.append(f"Column {col_letter(c)} looks like a rate but does not "
                         f"say which: choose Supply or Installation for it, or "
                         f"Ignore it.")
    seen = {}
    for c in grid["cols"]:
        t = mapping.get(str(c), "")
        if t and t != UNDECIDED:
            seen.setdefault(t, []).append(col_letter(c))
    for t, letters in seen.items():
        if len(letters) > 1:
            probs.append(f"{TARGET_LABEL[t]} is chosen for {len(letters)} columns "
                         f"({', '.join(letters)}). Keep one.")
    if "description" not in seen:
        probs.append("Choose the column that holds the Description.")
    if "qty" not in seen:
        probs.append("Choose the column that holds the Quantity. Only the total "
                     "quantity is imported; floor or area columns are left out.")
    return probs


# =============================================================================
# 6. THE ROWS — lines, headers, sections, totals, flags
# =============================================================================

_RATE_ONLY = re.compile(r"^\s*(?:r\s*\.?\s*o\s*\.?|i\s*\.?\s*r\s*\.?|r\s*/\s*o|rate\s*only)\s*$", re.I)
# A RATE-ONLY line (1 October 2026): exactly the four spellings the brief
# named — "RO", "R.O.", "R/O", "RATE ONLY" — any case, surrounding spaces
# ignored. Quantity 0, rates kept, never a flag on its own. Anything ELSE that
# `_RATE_ONLY` matches ("I.R.", "R. O.") keeps v1's rule: blank, red, blocked.
_RO_QTY = re.compile(r"^\s*(?:r\.?o\.?|r\s*/\s*o|rate\s+only)\s*$", re.I)
# What a rate-only line's remark says. Captured on the record, never printed.
RATE_ONLY_REMARK = "Rate only (RO) on the source sheet"
_NA = re.compile(r"^\s*(?:n\s*\.?\s*a\s*\.?|n\s*/\s*a|nil|-+|—|–)\s*$", re.I)
_DASH = re.compile(r"^\s*(?:-+|—|–)\s*$")
_ERROR_TEXT = re.compile(r"^\s*#(?:VALUE!|REF!|DIV/0!|N/A|NAME\?|NUM!|NULL!|SPILL!|CALC!)\s*$", re.I)
_NUMERIC_TEXT = re.compile(r"^\s*[-+]?(?:\d{1,3}(?:,\d{2,3})+|\d+)(?:\.\d+)?\s*$")
_ARITH = re.compile(r"^\s*\(?\s*\d+(?:\.\d+)?\s*\)?(?:\s*[-+*/xX×]\s*\(?\s*\d+(?:\.\d+)?\s*\)?)+\s*$")

_TOTAL_LABEL = re.compile(
    r"^\s*(?:grand\s*total|g\.?\s*total\b|sub\s*-?\s*total|total\b|"
    r"carried\s+(?:forward|over)|brought\s+forward|c\s*/\s*f\b|b\s*/\s*f\b)", re.I)
# ⚠ Widened 30 September 2026: "TOTAL AMOUNT   A+B+C" (the Jamnagar sheet's
#   row 124) is a grand total with no brackets round its sum.
_GRAND = re.compile(r"grand\s*total|\bg\.?\s*total\b|total\s*\([^)]*\+[^)]*\)|net\s+total|"
                    r"\btotal\b.*\b[a-z0-9]{1,3}\s*\+\s*[a-z0-9]{1,3}\b", re.I)
# A label that says total ANYWHERE — "BASIC VALUE SUBTOTAL (B) >>>>" on the
# Sify sheet. Read as a total only on a row with a figure and no rate.
_TOTAL_ANYWHERE = re.compile(r"\bsub\s*-?\s*total\b|\btotal\b", re.I)
_PARTIAL = re.compile(r"sub\s*-?\s*total|carried|brought|c\s*/\s*f|b\s*/\s*f|total\s+of\s+\S", re.I)
# A label that is NOTHING BUT a total word. Such a row is a total even with no
# figure on it (its formula may have no saved value); any longer label —
# "Total flooding system" is a heading — has to carry a figure somewhere on
# the row before it is treated as a total and dropped.
_BARE_TOTAL = re.compile(r"^\s*(?:grand\s*|g\.?\s*|sub\s*-?\s*|net\s*)?total\s*[:.\-=>]*\s*$", re.I)
TOTAL_LABEL_MAX = 60

# A section code in the item column: "A", "B", "II", "SECTION B", "Part C".
# The prefix word is case-insensitive; the code itself must be CAPITALS, so a
# sub-item "a" or "ii" is never read as a section.
_SECTION_ITEM = re.compile(r"^(?:(?i:section|part|schedule|system)\s*[-:.]?\s*)?([A-Z]{1,2}|[IVX]{1,4})\s*[.):]?$")
_SECTION_WORD = re.compile(r"^(?:section|part|schedule)\s*[-:.]?\s*([A-Za-z0-9]{1,3})\b[\s.:\-–—]*(.*)$", re.I | re.S)
_SUB_LABEL = re.compile(r"^\(?(?:[a-z]{1,2}|[ivx]{1,4})\)?\.?$")
# "(A) BOQ FOR HYDRANT SYSTEM" — a section heading carrying its own code.
_SECTION_PAREN = re.compile(r"^\(\s*([A-Z]{1,2})\s*\)\s*(.+)$", re.S)
# A sub-label the SHEET wrote at the start of a description: "a) 150 mm",
# "(b) 100 mm", "c. 80 mm". Lower-case letters only, so "A." or a size such
# as "80 mm" is never read as one.
_DESC_LABEL = re.compile(r"^\s*(?:\(([a-z]{1,2})\)|([a-z]{1,2})\)|([a-z])\.)\s+(\S.*)$", re.S)

# Flag kinds. RED ones on the QUANTITY leave it blank and BLOCK the save on the
# form until somebody types a figure or removes the line; see boqimport.py and
# ABOUT.md §7 gap 42 for why a blank quantity would otherwise be saved as 0.
_FLAG_TEXT = {
    "rate_only": "Rate only on source sheet — quantity left blank",
    "na":        "“{raw}” on the source sheet — left blank",
    "error":     "Excel error {raw} — left blank",
    "date":      "a date ({raw}) where a number belongs — left blank",
    "formula":   "formula, no saved value — open and save in Excel, then import again",
    "arith":     "sum written as text ({raw}) — not worked out; left blank",
    "text":      "text “{raw}” where a number belongs — left blank",
    "long":      "text longer than {n:,} characters — cut; paste the rest by hand",
    "no_item":   "a priced line with no item number and no numbered line above it — type one",
    "no_desc":   "figures but no description — type one",
    "item_date": "the item number is a date on the source sheet ({raw}) — left blank",
    "dup_section": "section {raw} appears again — this one is named {new}; rename it if you like",
    # 30 September 2026 — the blocking checks (they need somebody to type).
    "no_qty":    "a priced line with no quantity — type it",
    "no_rate_amt": "{track} amount {amount} but no {track} rate — type the rate",
    "no_rate":   "a quantity but no rate on either track — type a rate",
    "mismatch":  ("{track}: quantity {qty} × rate {rate} = {calc}, but the sheet's amount "
                  "is {amount} — the rate is left blank; type the right one"),
    # …and the ones that only ask for a look.
    "mismatch_both": ("quantity {qty} × (supply + installation rate) = {calc}, but the "
                      "sheet's amount is {amount} — review"),
    "lump_sum":  "an amount with no quantity or rate — imported as 1 LS at {amount}; review",
    "total_bad": "“{label}” does not add up — sheet {sheet}, the lines above add up to {calc}",
    # 1 October 2026 — the Iron Mountain sheet.
    # A rate-only line ("RO" in the quantity) with no rate on either track. It
    # is the ordinary blocking "rate" need, so "Not priced" answers it.
    "ro_no_rate": "rate-only line with no rate",
    # Notes only — listed under "other notes from the reader", never a need.
    "below_grand": "not imported — row {grand} is the grand total: {raw}",
    "make_number": "a number ({raw}) in the Make column — not kept",
    # 6 October 2026 (CLIENT_CHANGES.md §0, forty-fourth block, R3) — the
    # discount. The qty × rate check, extended: with a discount it is qty ×
    # the NET rate, and a sheet's own net rate is checked against rate less
    # discount. Both block, the rate left blank — v1's "mismatch" rule.
    "mismatch_net": ("{track}: quantity {qty} × net rate {net} ({rate} less {disc}%) = {calc}, "
                     "but the sheet's amount is {amount} — the rate is left blank; type the "
                     "right one"),
    "net_mismatch": ("{track}: rate {rate} less {disc}% = {calc}, but the sheet's net rate is "
                     "{net} — the rate is left blank; type the right one"),
    # …and a note: a discount that is no percentage is not read as one.
    "disc_range": ("a discount of {raw} is not a % from 0 to 100 — left blank; a discount "
                   "in rupees belongs in the Remark"),
    # 6 October 2026 (CLIENT_CHANGES.md §0, forty-fifth block) — the BOQ AS IT
    # IS. Every one of these is a NOTE: amber, never a need, never a block, and
    # no figure is changed by it.
    "as_remark": "“{raw}” is not a number — the box is left blank and the text kept in the remark",
    "as_mismatch": ("{track}: the sheet says {amount}, the BOQ computes {calc} (quantity {qty} × "
                    "net rate {net}) — the rate is kept as the sheet has it; the BOQ bills "
                    "quantity × rate"),
    "as_no_rate_amt": ("{track}: the sheet says {amount}, but the line has no {track_l} rate, "
                       "so the BOQ has no amount for it"),
    "as_net": ("{track}: rate {rate} less {disc}% is {calc}, the sheet's net rate is {net} — "
               "the rate and the discount are kept as the sheet has them"),
}

# The short name a numeric box goes by in a line's remark when the sheet wrote
# text there (A1): "Qty: Included", "Supply rate: #VALUE! in sheet".
REMARK_FIELD = {
    "qty": "Qty", "supply_base_rate": "Supply base rate",
    "escalation_pct": "Supply escalation %", "supply_rate": "Supply rate",
    "install_base_rate": "Installation base rate",
    "install_escalation_pct": "Installation escalation %",
    "install_rate": "Installation rate", "supply_amount": "Supply amount",
    "install_amount": "Installation amount", "amount": "Amount",
    "supply_disc_pct": "Supply disc %", "install_disc_pct": "Installation disc %",
    "supply_net_rate": "Supply net rate", "install_net_rate": "Installation net rate",
}


def _half_up(x: float) -> float:
    """`boq.round_half_up()`'s rule, restated because this leaf may import
    nothing of the app's (and only its listed standard library): the figure to
    15 significant digits, then a half rounds AWAY from zero — Excel's ROUND
    (A8). Worked on the digits, as `_BOQ_JS` `round2()` does, so binary noise
    cannot move it: 2.125 -> 2.13, 5.025 (5.0249… in binary) -> 5.03.
    `tests/test_boq_as_is.py` holds this, `boq.net_of()` and the editor equal."""
    x = float(x)
    neg, s = x < 0, f"{abs(x):.15g}"
    if "e" in s or "n" in s:
        return x                      # nowhere near a rate: left as it is
    ip, _dot, fp = s.partition(".")
    fp += "000"
    digits = list(ip + fp[:2])
    if fp[2] >= "5":
        k = len(digits) - 1
        while k >= 0 and digits[k] == "9":
            digits[k] = "0"
            k -= 1
        if k < 0:
            digits.insert(0, "1")
        else:
            digits[k] = str(int(digits[k]) + 1)
    whole = "".join(digits)
    v = float(whole[:-2] + "." + whole[-2:])
    return -v if neg else v

# How far a figure may be off and still agree: a rupee, the totals check's own
# tolerance. Used by the footing checks and the qty × rate check alike.
FOOT_TOLERANCE = TOTALS_TOLERANCE

# The fields a blocking flag can point the form at. "rate" means "either unit
# rate" and marks the supply one.
NEED_FIELDS = ("item_no", "description", "total_qty", "supply_rate", "install_rate", "rate")


def _fmt(v) -> str:
    """A figure for a flag sentence: Indian-agnostic, two decimals only when
    they are not zero. Never used for anything but the sentence."""
    f = float(v)
    return f"{f:,.0f}" if f.is_integer() else f"{f:,.2f}"


def sub_label(n: int) -> str:
    """0 -> "a", 25 -> "z", 26 -> "aa", 27 -> "ab" — the sub-item letters,
    continuing past z the way a spreadsheet's columns do."""
    return col_letter(n).lower()


def _label_index(label: str) -> int:
    """"a" -> 0, "z" -> 25, "aa" -> 26: the inverse of `sub_label`, or -1."""
    s = str(label or "").strip().lower().strip("().")
    if not s or not s.isalpha() or not s.isascii():
        return -1
    n = 0
    for ch in s:
        n = n * 26 + (ord(ch) - 96)
    return n - 1


def split_sheet_label(desc: str):
    """`(label, rest)` when a description opens with the sheet's own sub-label
    ("a) 150 mm" -> ("a", "150 mm")), else `("", desc)`."""
    m = _DESC_LABEL.match(desc or "")
    if not m:
        return "", desc
    return (m.group(1) or m.group(2) or m.group(3)), m.group(4).strip()


class Structure:
    """
    The item numbering the sheet implies — ABOUT.md §5 `/boq/import`, *The
    structure* (30 September 2026). Pure: rows go in, placements come out.

    Feed it one section at a time with `place(item, desc, priced)`; call
    `new_section()` at every section break. Each placement is a dict:

        {"kind": "subheading" | "header" | "item" | "spec_text" | "sub_item",
         "item_no", "parent_item_no", "is_header", "item_src", "description"}

    `item_src` is "sheet" (the sheet wrote the number or the sub-label),
    "auto" (the number was worked out from the rows above) or "". It lives
    on the import payload only and never reaches a record.

    The rules, walked in order down a section:
      1. Before the first NUMBERED row, an unpriced row is a sub-heading: no
         number, never flagged.
      2. A numbered row starts a new item and resets the sub-item counter —
         unpriced it is a header (the spec line), priced a standalone item.
         A number that is a child of an open header ("24.a" under "24",
         "2.1.4.1" under "2.1.4") or a bare sub-label ("a)") is that header's
         sub-item instead, exactly as v1 read it.
      3. An unnumbered, unpriced row after a numbered one is spec text under it.
      4. An unnumbered, priced row after a numbered one is its sub-item:
         <parent>.a, .b … .z, .aa, .ab — `auto`.
      5. A sub-label the sheet wrote at the start of the description ("a)",
         "(a)", "a.") is used as it stands — `sheet`.
      6. The counter resets at every numbered row and every section.

    A sub-item's description is a size label and is stripped; a clause (a
    header, a standalone item, spec text) is carried exactly as the sheet
    wrote it. That is the picker's own rule — `boq._seed_line()` strips a
    variant label and copies `spec_text` verbatim — and the save strips both.
    """

    def __init__(self):
        self.new_section()

    def new_section(self):
        self.parent = ""            # the current numbered row's item number
        self.parent_header = False  # …and whether it is an unpriced header
        self.heads = []             # every header number in this section
        self.sub_n = 0
        self.numbered = False       # a numbered row has been seen in this section
        self.placed = 0             # rows placed in this section

    def fresh(self) -> bool:
        """Nothing placed yet in this section."""
        return self.placed == 0

    def _bump(self, label: str):
        idx = _label_index(label)
        if idx >= 0:
            self.sub_n = max(self.sub_n, idx + 1)

    def place(self, item: str, desc: str, priced: bool) -> dict:
        self.placed += 1
        out = {"item_no": item, "parent_item_no": "", "is_header": not priced,
               "item_src": "sheet" if item else "", "description": desc}
        if item:
            parent = ""
            for h in self.heads:
                if _is_child(item, h) and len(h) > len(parent):
                    parent = h
            if not parent and self.parent and self.parent_header and _SUB_LABEL.match(item):
                parent = self.parent
            self.numbered = True
            if parent:
                out.update(kind="sub_item" if priced else "header", parent_item_no=parent)
                if priced:
                    out["description"] = (desc or "").strip()
                tail = item[len(parent):] if item.startswith(parent) else item
                self._bump(tail)
                if not priced:
                    self.heads.append(item)
                    self.parent, self.parent_header, self.sub_n = item, True, 0
                return out
            self.parent, self.parent_header, self.sub_n = item, not priced, 0
            if not priced:
                self.heads.append(item)
            out["kind"] = "header" if not priced else "item"
            return out

        if not self.numbered or not self.parent:
            out["kind"] = "subheading" if not priced else "item"
            return out
        if not priced:
            out.update(kind="spec_text", parent_item_no=self.parent)
            return out
        label, rest = split_sheet_label(desc)
        if label:
            self._bump(label)
            out.update(item_no=f"{self.parent}.{label}", description=rest, item_src="sheet")
        else:
            out.update(item_no=f"{self.parent}.{sub_label(self.sub_n)}", item_src="auto",
                       description=(desc or "").strip())
            self.sub_n += 1
        out.update(kind="sub_item", parent_item_no=self.parent)
        return out


class Footing:
    """
    Running sums per amount column, and the arithmetic that recognises a
    total row — 30 September 2026. A figure is checked, in order, against:
    the lines since the last subtotal (a SUBTOTAL); the subtotals of this
    section, or all of its lines (a SECTION TOTAL); every line so far (a
    GRAND TOTAL); and, for a single figure, supply + installation together
    (a COMBINED TOTAL, whichever column it sits in). Within ±1.00.
    """

    def __init__(self, cols: list):
        self.cols = list(cols)
        z = lambda: {c: 0.0 for c in self.cols}           # noqa: E731
        self.grand = z()
        self._z = z
        self.new_section()

    def new_section(self):
        self.run, self.subs, self.sec = self._z(), self._z(), self._z()
        self.run_n = self.subs_n = 0

    def add(self, amounts: dict):
        for c in self.cols:
            v = float(amounts.get(c) or 0.0)
            self.run[c] += v
            self.sec[c] += v
            self.grand[c] += v
        self.run_n += 1

    @staticmethod
    def _near(a, b) -> bool:
        return abs(float(a) - float(b)) <= FOOT_TOLERANCE

    def check(self, figures: dict, grand_label: bool = False) -> dict:
        """`{"kind", "status", "figures": [{"col", "sheet", "lines"}]}`.
        A row whose LABEL says grand total is tried as one first — on a
        one-section sheet the same figure is also that section's total."""
        figs = {c: float(v) for c, v in figures.items() if c in self.run}
        cands = [("grand total", self.grand)] if grand_label else []
        cands += [("subtotal", self.run)]
        if self.subs_n:
            cands.append(("section total", self.subs))
        cands += [("section total", self.sec), ("grand total", self.grand)]
        kind = ""
        shown = {c: (self.run if self.run_n else self.sec)[c] for c in figs}
        if figs:
            for k, base in cands:
                if all(self._near(v, base[c]) for c, v in figs.items()):
                    kind, shown = k, {c: base[c] for c in figs}
                    break
            if not kind and len(figs) == 1 and len(self.cols) >= 2:
                (c1, v), = figs.items()
                for base in (self.run, self.sec, self.grand):
                    if self._near(v, sum(base.values())):
                        kind, shown = "combined total", {c1: sum(base.values())}
                        break
        out = {"kind": kind or "total", "status": "match" if kind else "mismatch",
               "figures": [{"col": c, "sheet": round(v, 2), "lines": round(shown[c], 2)}
                           for c, v in figs.items()]}
        if kind == "subtotal":
            for c, v in figs.items():
                self.subs[c] += v
            self.subs_n += 1
        if kind:
            self.run, self.run_n = self._z(), 0
        return out


def _numeric(value, kind: str, field: str):
    """
    `(number | None, flag_kind | "")` for one quantity or rate cell.

    Blank is `(None, "")`. A number is itself — ×100 when it is a percentage
    in an escalation column, which is how the client's Sify sheet stores 15%
    and how `tools/gen_demo_data.py` read it. Everything else is `None` and a
    flag: this function never turns a word into a figure, and the only text it
    reads as a number is a plain number with thousands separators.
    """
    if kind == "formula":
        return None, "formula"
    if kind == "error":
        return None, "error"
    if kind == "date":
        return None, "date"
    if value is None:
        return None, ""
    if _is_number(value):
        v = float(value)
        if kind == "pct" and field in PCT_FIELDS:
            v = v * 100.0
        return v, ""
    s = str(value).strip()
    if not s:
        return None, ""
    if field == "qty" and _RO_QTY.match(s):
        # A rate-only line (1 October 2026): 0, and "ro" — a MARK, not a flag.
        # `build()` clears it before any flag is raised.
        return 0.0, "ro"
    if field == "qty" and _RATE_ONLY.match(s):
        return None, "rate_only"
    if field != "qty" and _DASH.match(s):
        # The client's "-" in a rate cell: negotiated directly, not escalated —
        # blank, and NOT a flag. `boq._opt_num()` reads it the same way.
        return None, ""
    if field in DISC_FIELDS:
        # A discount cell (6 October 2026, R3): "10%" is how people write one,
        # and "NA" / "nil" in it means no discount — blank, not a flag.
        if s.endswith("%"):
            s = s[:-1].strip()
        if _NA.match(s):
            return None, ""
    if _ERROR_TEXT.match(s):
        return None, "error"
    if _NA.match(s):
        return None, "na"
    if _NUMERIC_TEXT.match(s):
        return float(s.replace(",", "")), ""
    if _ARITH.match(s):
        return None, "arith"
    return None, "text"


def _item_text(value, kind: str) -> str:
    """
    An item number as text. Never a float.

    ⚠ **A numeric item is ROUNDED, then its trailing zeros are stripped** (1
      October 2026 — the Iron Mountain sheet, whose every item cell is
      `0.00`-formatted and holds 5.199999999999999 for item 5.2). Rounded to
      the decimals of its number format when that is a FIXED one ("0.00" → 2),
      otherwise to at most `ITEM_DECIMALS`: 5.199999999999999 → "5.2",
      17.200000000000003 → "17.2", 3.01 → "3.01", 4.0 → "4". A `0.00` cell
      holding 1.1 read "1.10" before this and reads "1.1" now. A TEXT item
      cell is exactly as the sheet wrote it, trimmed.
    """
    if value is None or kind in ("formula", "error", "date"):
        return ""
    if _is_number(value):
        places = int(kind[3:]) if kind.startswith("dp:") else ITEM_DECIMALS
        s = f"{float(value):.{places}f}"
        if "." in s:
            s = s.rstrip("0").rstrip(".")
        return "0" if s == "-0" else s
    return str(value).strip()


def _num_display(value, kind: str) -> str:
    """A number in a TEXT column (a description, a unit) as Excel shows it —
    v1's rule, unchanged: a fixed format keeps its decimals, otherwise ten
    significant digits. Item numbers have their own rule, `_item_text`."""
    if kind.startswith("dp:"):
        return f"{float(value):.{int(kind[3:])}f}"
    f = float(value)
    return str(int(f)) if f.is_integer() else f"{f:.10g}"


def _plain(value, kind: str) -> str:
    """Any other text cell: description, unit."""
    if value is None or kind == "formula":
        return ""
    if _is_number(value):
        return _num_display(value, kind)
    return str(value).strip()


def _is_caps_title(s: str) -> bool:
    letters = [ch for ch in s if ch.isalpha()]
    return (3 <= len(letters) and len(s) <= 120 and "\n" not in s
            and all(not ch.islower() for ch in letters))


def _section_of(item: str, desc: str):
    """(code, title) when this row is a section break, else None. Code "" means
    "assign the next free letter"."""
    m = _SECTION_ITEM.match(item)
    if item and m:
        return m.group(1), desc
    if not item:
        m = _SECTION_WORD.match(desc)
        if m:
            return m.group(1).upper(), (m.group(2).strip() or desc)
        m = _SECTION_PAREN.match(desc)
        if m and _is_caps_title(desc):
            return m.group(1), m.group(2).strip()
        if desc and _is_caps_title(desc):
            return "", desc
    return None


def _is_child(item: str, head: str) -> bool:
    if not head or item == head or not item.startswith(head):
        return False
    nxt = item[len(head)]
    return nxt in ".-/ (" or (head[-1].isdigit() and nxt.isalpha())


def _repeats_header(vals: list, hdr_labels: dict) -> bool:
    """A later row that repeats the column heading — the Sify sheet prints it
    again above section B. Two or more cells must share a keyword with the
    heading's own label in the same column."""
    same = 0
    for ci, v in enumerate(vals):
        if isinstance(v, str) and v.strip() and ci in hdr_labels:
            if _hits(v) & _hits(hdr_labels[ci]):
                same += 1
    return same >= 2


def _heading_remnant(vals: list, item_ci, desc_ci, fig_cis: set) -> bool:
    """The rest of a heading the header detector did not take: no figure
    anywhere, no item number or description, and what text sits in the
    quantity, rate and amount columns is heading words ("U/ Rate", "AMT")."""
    if any(_is_number(v) for v in vals):
        return False
    for ci in (item_ci, desc_ci):
        if ci is not None and ci < len(vals) and isinstance(vals[ci], str) and vals[ci].strip():
            return False
    words = [vals[ci] for ci in fig_cis if ci < len(vals)
             and isinstance(vals[ci], str) and vals[ci].strip()]
    return bool(words) and all(_hits(w) for w in words)


def _heading_words_only(vals: list) -> bool:
    """No figure, and every label a keyword, "total" or a short code — the
    lower half of a two-row heading ("Rate | Amount")."""
    seen = False
    for v in vals:
        if _is_number(v):
            return False
        if isinstance(v, str) and v.strip():
            t = _norm(v)
            if not (_hits(t) or _TOTAL_WORD.search(t) or len(t) <= 3):
                return False
            seen = True
    return seen


def _row_summary(vals: list, limit: int = 4) -> str:
    """What a row holds, for a note: its first few non-empty cells — text in
    quotes, cut at 60 characters, and non-zero figures. "" for a row with
    nothing on it but blanks and zeros."""
    bits = []
    for v in vals:
        if _is_number(v):
            if abs(v) >= 0.005:
                bits.append(_fmt(v))
        elif isinstance(v, str) and v.strip():
            t = re.sub(r"\s+", " ", v).strip()
            bits.append(f"“{t[:57]}…”" if len(t) > 60 else f"“{t}”")
        if len(bits) > limit:
            bits[limit:] = ["…"]
            break
    return " · ".join(bits)


def _group_labels(lines: list) -> int:
    """
    Group labels inside one item (1 October 2026 — the Iron Mountain sheet's
    sluice valves: "PN-25" over DN 250 … DN 80, then "PN-16" over the same
    sizes again). Returns how many rows became labels.

    The rule, exactly: under ONE parent (the line the rows' `parent_item_no`
    names, by THE parent rule — same section, nearest above), when TWO OR MORE
    spec-text rows are each IMMEDIATELY followed by a child of that parent,
    those rows are group labels. Each label goes in front of the description
    of every child after it, up to the next label or the first line that is
    not this parent's child: "PN-25 · DN 250". The label row becomes
    `kind: "group_label"`, and `boqimport.editor_model()` does NOT append it to
    the parent's text.

    ⚠ ONE such row, or rows that are not each followed by a child, stay spec
      text exactly as before — the deluge valve's "i) Mains deluge valve" /
      "ii) Pressure gauges…" (only the second is followed by a child) and a
      monitor's run of specification bullets. "Immediately followed" is the
      next LINE: a blank row between a label and its first size is no line.
    """
    def parent_at(i):
        l = lines[i]
        for k in range(i - 1, -1, -1):
            o = lines[k]
            if o["section"] != l["section"]:
                return None
            if o["item_no"] and o["item_no"] == l["parent_item_no"]:
                return k
        return None

    def child_of(o, l) -> bool:
        return (o["section"] == l["section"] and o["kind"] in ("sub_item", "header")
                and o["parent_item_no"] == l["parent_item_no"])

    cands = {}
    for i, l in enumerate(lines):
        if l["kind"] != "spec_text" or not l["parent_item_no"]:
            continue
        if i + 1 < len(lines) and child_of(lines[i + 1], l):
            p = parent_at(i)
            if p is not None:
                cands.setdefault(p, []).append(i)

    n = 0
    for idxs in cands.values():
        if len(idxs) < 2:
            continue
        for i in idxs:
            lines[i]["kind"] = "group_label"
            n += 1
        for i in idxs:
            l = lines[i]
            label = (l["description"] or "").strip()
            for o in lines[i + 1:]:
                if (o["kind"] == "group_label" or o["section"] != l["section"]
                        or o["parent_item_no"] != l["parent_item_no"]):
                    break
                if child_of(o, l):
                    o["description"] = f"{label} · {(o['description'] or '').strip()}"
                    o["group_label"] = label
    return n


def build(grid: dict, mapping: dict, guided: bool = False) -> dict:
    """
    The confirmed mapping applied to every row under the header.

    ⚠ **TWO MODES from 6 October 2026** (CLIENT_CHANGES.md §0, forty-fifth
      block — the client: *"open my sheet as a BOQ AS IT IS"*).

      * **As it is — the default, and the BOQ's** (`guided=False`, A1/A2).
        Every mapped cell comes through as the sheet has it: blank → `None`,
        0 → 0, a number → itself. TEXT in a numeric column ("Included",
        "By client", "NA", "Nil", a sum written as text, a date, TRUE) leaves
        the figure `None` and goes into the line's remark as `<Field>: <text>`
        (`line["as_remark"]`); an Excel error as `<Field>: #VALUE! in sheet`;
        a bare "-" is just blank. A formula's cached value is what is read,
        and a cached "" — or none at all, which openpyxl cannot tell apart — is
        blank (`counts["formula_blank"]`). **Nothing is a need and nothing
        blocks**: no "no quantity", no "no rate", no "no item number", no "no
        description"; a quantity × rate that differs from the sheet's amount
        KEEPS the rate and is an amber note giving both figures
        (`as_mismatch`). The kept rules: item-number rounding, a %-formatted
        cell read as the % Excel shows, "RO" as a rate-only line at 0, a lump
        sum for an amount alone, the footing checks and the row
        classification.
      * **Guided** (`guided=True`) — the 30 September / 1 October 2026 rules,
        exactly as they were: blocking `needs`, a red `block` on a quantity the
        reader could not take, a mismatched rate left BLANK. **The work-order
        importer passes it** (`workorder.py`): the forty-fifth block is about
        the BOQ, and the work order's own rulings (D, R1) are built on these
        needs.
    """
    as_is = not guided
    return _build(grid, mapping, as_is)


def _build(grid: dict, mapping: dict, as_is: bool) -> dict:
    """
    `build()`'s body.

    Returns::

        {"sections": [{"code", "title"}],
         "lines":    [{"row", "kind", "item_no", "parent_item_no", "item_src",
                       "section", "is_header", "description", "unit", "make",
                       "qty", <rate fields>, "lump_sum", "rate_only",
                       "flags": [...], "needs": [{"field", "message"}],
                       "block": bool, ["group_label": str]}],
                     # kind: "subheading" | "header" | "item" | "spec_text" |
                     #       "sub_item" | "group_label" (1 October 2026)
         "flags":    [{"row", "col", "field", "raw", "kind", "message",
                       "severity"}],
         "needs":    [{"row", "field", "message"}],     # every blocking flag
         "checks":   [{"row", "label", "kind", "status", "figures"}],
         "counts":   {"lines", "headers", "sections", "totals_dropped",
                      "flagged", "auto_items", "sheet_items", "lump_sums",
                      "needs", "repeats", "rate_only", "rate_only_no_rate",
                      "below_grand", "group_labels"},
         "totals":   {...}}

    Numbers are floats or None — shaping them for the form is `boqimport.py`'s
    job. `block` is True on a line whose QUANTITY was left blank by a flag.

    30 September 2026 — the structure, the totals and the real flags:

    * **Item numbers come from the sheet's own shape** — `Structure`. A row
      the sheet did not number is spec text, a sub-heading, or a sub-item
      <parent>.a, and is no longer flagged "no item number".
    * **An amount with no quantity and no rate** is checked by arithmetic
      (`Footing`): a subtotal, a section total, a grand total or a combined
      total is a FOOTING CHECK and never a line; anything else is a LUMP SUM
      — 1 LS at the amount, flagged for review. **A zero amount is no
      amount**: such a row is read as if the amount cell were empty.
    * **A row labelled as a total** is checked the same way, and flagged with
      both figures when it does not add up. It is never a line.
    * **A later row repeating the column heading** is skipped, never a line
      and never flagged.
    * **Blocking flags** (`needs`) — a priced line with no quantity; an amount
      or a quantity with no rate; quantity × rate more than ±1.00 from the
      sheet's amount (the rate is then left BLANK, both figures in the flag);
      and a priced line with no item number or no description, which the
      save would refuse anyway. They are the fields the form leads to.

    1 October 2026 — the Iron Mountain sheet (CLIENT_CHANGES.md §0,
    thirty-ninth block):

    * **A rate-only line** — "RO", "R.O.", "R/O" or "Rate only" in the
      quantity — is quantity 0 with its rates kept and `rate_only` set. That
      alone is not a flag. With no rate on either track (a zero is none, the
      quantity being no real one) it gets the ordinary blocking "rate" need,
      worded "rate-only line with no rate", so *Not priced* answers it.
    * **Below the grand total.** Once a row is recognised as the grand total
      by BOTH its label and its sums, no later row becomes a line: each is
      listed as a "below_grand" note. A later total row is still a footing
      check, which is not a line (the Jamnagar sheet's row 125 sits under its
      grand total). Rows above it are read exactly as before.
    * **A number in the Make column** is not a make: dropped, and noted.
    * **Group labels inside an item** — `_group_labels()`, after the walk.
    """
    cols = grid["cols"]
    pos = {str(c): i for i, c in enumerate(cols)}
    col_of = {}
    for c in cols:
        t = mapping.get(str(c), "")
        if t and t != UNDECIDED and t not in col_of:
            col_of[t] = pos[str(c)]
    mapped_cis = set(col_of.values())
    num_cis = {col_of[f] for f in ("qty",) + RATE_FIELDS if f in col_of}
    fig_cis = num_cis | {col_of[f] for f in AMOUNT_FIELDS if f in col_of}

    tracks = [t for t in ("supply", "install") if f"{t}_rate" in col_of]
    amount_cols = [f for f in AMOUNT_FIELDS if f in col_of]

    # ── 6 October 2026 (CLIENT_CHANGES.md §0, forty-fourth block, R1/R2) ────
    # Only the TICKED columns are read for a line's data, flags and needs: a
    # column mapped to nothing is ignored completely. (Whether a ROW is a total
    # or a repeated heading is still decided from the whole row — that is the
    # sheet's shape, not a column's data.)
    ticked_cis = sorted(pos[str(c)] for c in cols
                        if mapping.get(str(c), "") not in ("", UNDECIDED))
    # No Item No. column ticked: the reader numbers the lines itself (R2 — a
    # field whose column was not ticked is never asked for), by the sheet's
    # own structure — a heading row a header, the priced rows under it its
    # sub-items, a priced row with no heading above it an item of its own.
    item_mapped = "item_no" in col_of
    # Is any rate, base, amount or net column ticked at all? With none, no
    # line is asked for a rate (R2): the user chose not to import any.
    rate_mapped = any(f in col_of for f in RATE_FIELDS + AMOUNT_FIELDS + NET_FIELDS)
    auto_no = {"n": 0}

    def track_of(col: str) -> str:
        if col == "supply_amount":
            return "supply"
        if col == "install_amount":
            return "install"
        return tracks[0] if len(tracks) == 1 else "both"

    flags, lines, sections, totals_rows, checks, needs = [], [], [], [], [], []
    counts = {"lines": 0, "headers": 0, "sections": 0, "totals_dropped": 0, "flagged": 0,
              "auto_items": 0, "sheet_items": 0, "lump_sums": 0, "needs": 0, "repeats": 0,
              "rate_only": 0, "rate_only_no_rate": 0, "below_grand": 0,
              "group_labels": 0,
              # As it is (6 October 2026): cells of text kept in a remark, and
              # formula cells with no saved value that came in blank.
              "as_remark": 0, "formula_blank": 0}

    def flag(rnum, field, kind, raw="", severity="amber", field_label=None, **extra):
        msg = _FLAG_TEXT[kind].format(raw=raw, n=MAX_CELL_CHARS, **extra)
        ci = col_of.get("supply_rate" if field == "rate" else field)
        f = {"row": rnum, "col": col_letter(cols[ci]) if ci is not None else "",
             "field": field_label or TARGET_LABEL.get(field, "Rate" if field == "rate" else field),
             "raw": raw, "kind": kind, "message": msg, "severity": severity}
        flags.append(f)
        return f

    # "\x00" is the placeholder for the section lines sit in before the sheet
    # names one; it is replaced by a real code once every code is known.
    DEFAULT = "\x00"
    cur_section = DEFAULT
    struct = Structure()
    foot = Footing(amount_cols)
    hdr = grid.get("header") or []
    hdr_labels = header_labels(grid) if hdr else {}
    prev_repeat = False
    # The row the grand total was recognised on, by its label AND its sums.
    # Nothing below it becomes a line (1 October 2026).
    grand_row = None

    def footed_alone(rnum, label, amounts_nz, nums) -> bool:
        """An amount and nothing else that adds up as a total: a footing check,
        recorded, never a line. False when it is no total by arithmetic."""
        chk = foot.check(amounts_nz)
        if chk["status"] != "match":
            return False
        chk.update(row=rnum, label=label)
        checks.append(chk)
        totals_rows.append({"row": rnum, "label": label,
                            "amounts": {f: nums[f] for f in AMOUNT_FIELDS}})
        counts["totals_dropped"] += 1
        return True

    for ri in range(data_start(grid), len(grid["rows"])):
        rnum, vals = grid["rows"][ri]

        # ── A repeated column heading: skipped, never a line, never flagged ──
        # A 0 in a quantity or rate column does not make a row data: the Sify
        # sheet leaves one on its repeated heading.
        if hdr and not any(_is_number(_value(grid, ri, ci)) and _value(grid, ri, ci) != 0
                           for ci in num_cis):
            if _repeats_header(vals, hdr_labels) or _heading_remnant(
                    vals, col_of.get("item_no"), col_of.get("description"), fig_cis):
                counts["repeats"] += 1
                prev_repeat = True
                continue
            if prev_repeat and _heading_words_only(vals):
                counts["repeats"] += 1
                prev_repeat = False
                continue
        prev_repeat = False

        def cell(field):
            ci = col_of.get(field)
            return _value(grid, ri, ci), _kind(grid, ri, ci)

        iv, ik = cell("item_no")
        item = _item_text(iv, ik)
        dv, dk = cell("description")
        desc = _plain(dv, dk)
        # Classified on the stripped text, carried VERBATIM — the seeded Sify
        # schedule keeps the workbook's trailing spaces, and the save strips.
        desc_raw = dv if (isinstance(dv, str) and desc and dk != "formula") else desc
        uv, uk = cell("unit")
        unit = _plain(uv, uk)
        mv, mk = cell("make")
        make = _plain(mv, mk)
        # The sheet's own notes column (6 October 2026, R1) — word for word.
        rv, rk = cell("remark")
        remark = _plain(rv, rk)

        nums, nflags, raws = {}, {}, {}
        for field in ("qty",) + RATE_FIELDS + AMOUNT_FIELDS + DISC_FIELDS + NET_FIELDS:
            v, k = cell(field)
            n, fk = _numeric(v, k, field)
            nums[field], nflags[field] = n, fk
            raws[field] = "" if v is None else str(v)

        # ── As it is (6 October 2026, A1) ────────────────────────────────────
        # A cell that is not a number keeps its WORDS, in the line's remark,
        # and its figure is blank. Its flag becomes "as_remark" — still a
        # flag, so the row is still a line the way it always was ("NA" in the
        # quantity of a described row is a line with no quantity) — but never
        # a need and never a block. A bare "-" and a formula with no saved
        # value are simply blank: nothing to keep.
        as_notes, as_flags, as_formula = [], [], 0
        if as_is:
            for field in ("qty",) + RATE_FIELDS + DISC_FIELDS + AMOUNT_FIELDS + NET_FIELDS:
                fk = nflags[field]
                if not fk or fk == "ro" or field not in col_of:
                    continue
                raw = raws[field].strip()
                if fk == "formula":
                    nflags[field] = ""
                    as_formula += 1
                    continue
                if _DASH.match(raw):
                    nflags[field] = ""
                    continue
                as_notes.append(f"{REMARK_FIELD[field]}: "
                                + (f"{raw} in sheet" if fk == "error" else raw))
                as_flags.append((field, raw))
                nflags[field] = "as_remark"
            # "NA" / "nil" in a discount column is no discount (the forty-
            # fourth block) — and, as it is, the words are kept too (A1).
            for f in DISC_FIELDS:
                raw = raws[f].strip()
                if (f in col_of and nums[f] is None and not nflags[f] and raw
                        and not _DASH.match(raw)):
                    as_notes.append(f"{REMARK_FIELD[f]}: {raw}")
                    as_flags.append((f, raw))
            # A discount that is no percentage (above 100, below 0) is not
            # read as one — and its figure is not lost: it goes to the remark.
            for f in DISC_FIELDS:
                d = nums[f]
                if d is not None and not (0.0 <= d <= 100.0) and f in col_of:
                    as_notes.append(f"{REMARK_FIELD[f]}: {raws[f].strip()}")
                    as_flags.append((f, raws[f].strip()))
                    nums[f] = None
        # The figures as the sheet wrote them, BEFORE the classification below
        # reads a 0 rate on a row with no quantity as no rate (the owner's 30
        # September ruling — a rule about what a row IS). As it is, a line
        # keeps the sheet's 0 (A1).
        sheet_nums = dict(nums)
        amounts = {f: nums[f] for f in amount_cols if nums[f] is not None}
        # ⚠ A ZERO amount is NO amount (the owner's ruling, 30 Sep 2026): the
        #   Jamnagar sheet carries a 0 in both amount cells of every spec row.
        amounts_nz = {f: v for f, v in amounts.items() if abs(v) >= 0.005}

        # A RATE-ONLY line (1 October 2026): "RO" in the quantity. Its 0 is
        # not a quantity anybody measured, so it does not make a zero rate
        # count as a rate — and "ro" is a mark, cleared before any flag.
        rate_only = nflags["qty"] == "ro"
        if rate_only:
            nflags["qty"] = ""
        has_qty = nums["qty"] is not None and not rate_only
        # ⚠ A ZERO rate on a row with no quantity is no rate, for the same
        #   reason as a zero amount: the Sify sheet leaves a 0 in its
        #   installation-rate column on subtotal and section rows.
        has_rate = any(nums[f] is not None and (has_qty or abs(nums[f]) >= 0.005)
                       for f in RATE_FIELDS)
        if not has_qty and not has_rate:
            for f in RATE_FIELDS:
                nums[f] = None if nums[f] == 0 else nums[f]

        # ── Total rows: a footing check, never a line ────────────────────────
        labels = [x for x in (item, desc, unit) if x]
        for ci, v in enumerate(vals):
            if ci not in mapped_cis and isinstance(v, str) and v.strip():
                labels.append(v.strip())
        qty_ci, item_ci = col_of.get("qty"), col_of.get("item_no")
        has_figure = any(
            (_is_number(v) and ci not in (qty_ci, item_ci)) or _kind(grid, ri, ci) == "formula"
            for ci, v in enumerate(vals))
        total_label = next((l for l in labels if len(l) <= TOTAL_LABEL_MAX
                            and _TOTAL_LABEL.match(l)), "")
        if not total_label and has_figure and not has_rate:
            total_label = next((l for l in labels if len(l) <= TOTAL_LABEL_MAX
                                and _TOTAL_ANYWHERE.search(l)), "")
        if (total_label and not has_qty and not nflags["qty"] and not rate_only
                and (has_figure or _BARE_TOTAL.match(total_label))):
            totals_rows.append({"row": rnum, "label": total_label,
                                "amounts": {f: nums[f] for f in AMOUNT_FIELDS}})
            counts["totals_dropped"] += 1
            if amounts:
                grand_label = bool(_GRAND.search(total_label))
                chk = foot.check(amounts, grand_label=grand_label)
                chk.update(row=rnum, label=total_label)
                checks.append(chk)
                if chk["status"] != "match":
                    worst = chk["figures"][0]
                    for fg in chk["figures"]:
                        if not Footing._near(fg["sheet"], fg["lines"]):
                            worst = fg
                            break
                    flag(rnum, worst["col"], "total_bad", label=total_label,
                         sheet=_fmt(worst["sheet"]), calc=_fmt(worst["lines"]))
                elif grand_row is None and grand_label and chk["kind"] == "grand total":
                    # Recognised by BOTH its label and its sums: the sheet's
                    # schedule ends here (1 October 2026).
                    grand_row = rnum
            continue

        has_num_flags = bool(nflags["qty"]) or any(nflags[f] for f in RATE_FIELDS)
        amount_only = (bool(amounts_nz) and not has_qty and not has_rate
                       and not has_num_flags and not rate_only)

        # ── Below the grand total: never a line (1 October 2026) ─────────────
        # A later total row was checked just above, like any other. Anything
        # else — an add-on, a running total, a declarations block — is listed
        # with its row number and goes no further. Above the grand total
        # nothing here runs, so those rows are read exactly as before.
        if grand_row is not None:
            if amount_only and footed_alone(rnum, desc or item, amounts_nz, nums):
                continue
            # The note quotes the TICKED columns only (R1): an unticked
            # column is never read into anything the user is shown.
            text = _row_summary([vals[ci] for ci in ticked_cis if ci < len(vals)])
            if text:
                flag(rnum, "", "below_grand", raw=text,
                     field_label="Below the grand total", grand=grand_row)
                counts["below_grand"] += 1
            continue

        if not (item or desc or has_qty or has_rate or has_num_flags or amount_only
                or rate_only):
            continue

        # ── An amount and nothing else: a total by arithmetic, or a lump sum ─
        if amount_only and footed_alone(rnum, desc or item, amounts_nz, nums):
            continue

        priced = has_qty or has_rate or has_num_flags or amount_only or rate_only

        # ── Section rows ────────────────────────────────────────────────────
        if not priced:
            sec = _section_of(item, desc)
            # A code-less capitals title straight under a section heading is
            # that section's sub-heading ("SPRINKLER SYSTEM" under "(B) BOQ
            # FOR SPRINKLER SYSTEM"), not a section of its own.
            if sec is not None and not (sec[0] == "" and cur_section != DEFAULT
                                        and struct.fresh()):
                code, title = sec
                sections.append({"code": code, "title": title, "row": rnum})
                cur_section = len(sections) - 1      # resolved to a code below
                struct.new_section()
                foot.new_section()
                auto_no["n"] = 0
                counts["sections"] += 1
                continue

        if item and not (desc or priced):
            continue                                  # an item number and nothing else

        # A NUMBER in the Make column is not a make (1 October 2026 — the Iron
        # Mountain sheet carries 80832, 49831.2 … there): dropped, and noted.
        if make and (_is_number(mv) or _NUMERIC_TEXT.match(make)):
            flag(rnum, "make", "make_number", raw=make, field_label="Make")
            make = ""

        # No Item No. column ticked (R2): the number the sheet's shape implies.
        numbered_here = False
        if not item_mapped and not item and (not priced or not struct.parent_header):
            auto_no["n"] += 1
            item, numbered_here = str(auto_no["n"]), True

        placed = struct.place(item, desc_raw, priced)
        if numbered_here:
            placed["item_src"] = "auto"
        line = {"row": rnum, "kind": placed["kind"], "item_no": placed["item_no"],
                "parent_item_no": placed["parent_item_no"], "item_src": placed["item_src"],
                "section": cur_section, "is_header": placed["is_header"],
                "description": placed["description"], "unit": unit, "make": make,
                "remark": remark,
                "qty": None, "lump_sum": False, "rate_only": False,
                "flags": [], "needs": [], "block": False}
        for f in RATE_FIELDS + DISC_FIELDS:
            line[f] = None
        if as_is:
            # A1: the words of every cell that was not a number, for the
            # line's remark (`boqimport.editor_model()` appends them), and one
            # amber note each for the preview — never on the line's own chip.
            line["as_remark"] = as_notes
            for f, raw in as_flags:
                flag(rnum, f, "as_remark", raw=raw)
            counts["as_remark"] += len(as_notes)
            counts["formula_blank"] += as_formula

        if ik == "date" and iv is not None:
            line["flags"].append(flag(rnum, "item_no", "item_date", raw=str(iv))["message"])
        if dk == "long":
            line["flags"].append(flag(rnum, "description", "long")["message"])

        def need(field, message):
            line["needs"].append({"field": field, "message": message})
            needs.append({"row": rnum, "field": field, "message": message})

        if line["is_header"]:
            counts["headers"] += 1
        else:
            line["qty"] = nums["qty"]
            # As it is (A1), the sheet's own figures — a 0 rate on a row with
            # no quantity stays 0 on the line (`sheet_nums`); guided, v1's.
            src = sheet_nums if as_is else nums
            for f in RATE_FIELDS:
                line[f] = src[f]
            # The discount per track (6 October 2026, R3): a % from 0 to 100.
            # Anything else — a rupee figure above 100, a negative — is NOT a
            # discount: left blank and noted below, never guessed into one.
            disc_bad = []
            for f in DISC_FIELDS:
                d = nums[f]
                if d is not None and not (0.0 <= d <= 100.0):
                    disc_bad.append(f)
                    d = None
                line[f] = d

            def net_of(t):
                """quantity × THIS is what the sheet should show: the rate less
                the discount, unrounded — the checks allow a rupee anyway."""
                r = line[f"{t}_rate"]
                if r is None:
                    return None
                d = line.get(f"{t}_disc_pct")
                return r * (100.0 - d) / 100.0 if d else r

            if rate_only:
                line["rate_only"] = True
                counts["rate_only"] += 1
            if amount_only:
                # A LUMP SUM: 1 LS at the amount, in the track(s) its amount
                # column belongs to. Soft — it asks for a look, not a figure.
                line.update(qty=1.0, unit="LS", lump_sum=True)
                for col, v in amounts_nz.items():
                    t = track_of(col)
                    rate_f = "install_rate" if t == "install" else "supply_rate"
                    line[rate_f] = (line[rate_f] or 0.0) + v
                line["flags"].append(flag(rnum, "description", "lump_sum",
                                          amount=_fmt(sum(amounts_nz.values())))["message"])
                counts["lump_sums"] += 1

            if as_is:
                # ── As it is (6 October 2026, A2/A3): notes, never needs ──────
                # Nothing here blanks, alters or fills a figure, and nothing is
                # asked for: a blank quantity, a blank rate, a blank item number
                # or description is what the sheet wrote. Where the sheet's own
                # amount (or net rate) is not what the BOQ will compute from the
                # line, an AMBER note gives both figures — the rate stays as the
                # sheet has it, and the BOQ's amount stays quantity × net rate,
                # because an RA bill bills quantity × rate.
                q = line["qty"]

                def boq_net(t):
                    r = line[f"{t}_rate"]
                    if r is None:
                        return None
                    d = line.get(f"{t}_disc_pct")
                    return _half_up(r * (100.0 - d) / 100.0) if d else r

                if not amount_only:
                    for t in ("supply", "install"):
                        rate_f = f"{t}_rate"
                        col = (f"{t}_amount" if f"{t}_amount" in amounts else
                               "amount" if "amount" in amounts and track_of("amount") == t else "")
                        amt = amounts.get(col) if col else None
                        rate = line[rate_f]
                        name = "Supply" if t == "supply" else "Installation"
                        disc = line.get(f"{t}_disc_pct")
                        net = boq_net(t)
                        net_sheet = nums.get(f"{t}_net_rate")
                        if rate is None and amt is not None and abs(amt) >= 0.005:
                            line["flags"].append(flag(
                                rnum, rate_f, "as_no_rate_amt", track=name,
                                track_l=name.lower(), amount=_fmt(amt))["message"])
                        elif rate is not None and net_sheet is not None                                 and not Footing._near(net, net_sheet):
                            line["flags"].append(flag(
                                rnum, rate_f, "as_net", track=name, rate=_fmt(rate),
                                disc=_fmt(disc or 0), calc=_fmt(net),
                                net=_fmt(net_sheet))["message"])
                        elif rate is not None and amt is not None and q is not None                                 and not Footing._near(q * net, amt):
                            line["flags"].append(flag(
                                rnum, rate_f, "as_mismatch", track=name, qty=_fmt(q),
                                net=_fmt(net), calc=_fmt(q * net),
                                amount=_fmt(amt))["message"])
                    if "amount" in amounts and track_of("amount") == "both":
                        amt = amounts["amount"]
                        sn, inn = boq_net("supply"), boq_net("install")
                        if sn is None and inn is None and abs(amt) >= 0.005:
                            line["flags"].append(flag(
                                rnum, "rate", "as_no_rate_amt", track="Combined",
                                track_l="supply or installation",
                                amount=_fmt(amt))["message"])
                        elif q is not None and (sn is not None or inn is not None)                                 and not Footing._near(q * ((sn or 0) + (inn or 0)), amt):
                            line["flags"].append(flag(
                                rnum, "amount", "mismatch_both", qty=_fmt(q),
                                calc=_fmt(q * ((sn or 0) + (inn or 0))),
                                amount=_fmt(amt))["message"])
                if rate_only and line["supply_rate"] is None and line["install_rate"] is None:
                    counts["rate_only_no_rate"] += 1
            else:
                for field in ("qty",) + RATE_FIELDS:
                    fk = nflags[field]
                    if not fk:
                        continue
                    red = field == "qty"
                    f = flag(rnum, field, fk, raw=raws[field],
                             severity="red" if red else "amber")
                    line["flags"].append(f"{TARGET_LABEL[field]}: {f['message']}")
                    if red:
                        line["block"] = True

                # ── The blocking checks ────────────────────────────────────────
                q = line["qty"]
                if q is None:
                    if line["block"]:
                        need("total_qty", line["flags"][-1] if line["flags"] else _FLAG_TEXT["no_qty"])
                    else:
                        msg = flag(rnum, "qty", "no_qty", severity="red")["message"]
                        line["flags"].append(msg)
                        need("total_qty", msg)
                rate_needed = False
                if not amount_only:
                    for t in ("supply", "install"):
                        rate_f = f"{t}_rate"
                        col = (f"{t}_amount" if f"{t}_amount" in amounts else
                               "amount" if "amount" in amounts and track_of("amount") == t else "")
                        amt = amounts.get(col) if col else None
                        rate = line[rate_f]
                        name = "Supply" if t == "supply" else "Installation"
                        disc = line.get(f"{t}_disc_pct")
                        net = net_of(t)
                        net_sheet = nums.get(f"{t}_net_rate")
                        if rate is None and amt is not None and abs(amt) >= 0.005:
                            msg = flag(rnum, rate_f, "no_rate_amt", severity="red",
                                       track=name, amount=_fmt(amt))["message"]
                            line["flags"].append(msg)
                            need(rate_f, msg)
                            rate_needed = True
                        elif rate is not None and net_sheet is not None \
                                and not Footing._near(net, net_sheet):
                            # The sheet's own net rate against rate less discount
                            # (6 October 2026, R3) — both figures, the rate blank.
                            msg = flag(rnum, rate_f, "net_mismatch", severity="red", track=name,
                                       rate=_fmt(rate), disc=_fmt(disc or 0), calc=_fmt(net),
                                       net=_fmt(net_sheet))["message"]
                            line["flags"].append(msg)
                            line[rate_f] = None          # never guessed: left blank
                            need(rate_f, msg)
                            rate_needed = True
                        elif rate is not None and amt is not None and q is not None \
                                and not Footing._near(q * net, amt):
                            if disc:
                                # quantity × the NET rate (R3, the check extended).
                                msg = flag(rnum, rate_f, "mismatch_net", severity="red",
                                           track=name, qty=_fmt(q), net=_fmt(net),
                                           rate=_fmt(rate), disc=_fmt(disc), calc=_fmt(q * net),
                                           amount=_fmt(amt))["message"]
                            else:
                                msg = flag(rnum, rate_f, "mismatch", severity="red", track=name,
                                           qty=_fmt(q), rate=_fmt(rate), calc=_fmt(q * rate),
                                           amount=_fmt(amt))["message"]
                            line["flags"].append(msg)
                            line[rate_f] = None          # never guessed: left blank
                            need(rate_f, msg)
                            rate_needed = True
                    if "amount" in amounts and track_of("amount") == "both":
                        amt = amounts["amount"]
                        s, i = line["supply_rate"], line["install_rate"]
                        sn, inn = net_of("supply"), net_of("install")
                        if s is None and i is None and abs(amt) >= 0.005:
                            msg = flag(rnum, "rate", "no_rate_amt", severity="red",
                                       track="Combined", amount=_fmt(amt))["message"]
                            line["flags"].append(msg)
                            need("rate", msg)
                            rate_needed = True
                        elif q is not None and (s is not None or i is not None) \
                                and not Footing._near(q * ((sn or 0) + (inn or 0)), amt):
                            line["flags"].append(flag(
                                rnum, "amount", "mismatch_both", qty=_fmt(q),
                                calc=_fmt(q * ((sn or 0) + (inn or 0))), amount=_fmt(amt))["message"])
                    # ⚠ A BASE rate prices a line too (6 October 2026, R2): the form
                    #   works the unit rate out from base + escalation, so a line
                    #   carrying only a base is not "no rate". And with no rate,
                    #   base, amount or net column ticked at all, no rate is asked
                    #   for — the user chose not to import one.
                    no_rate = (line["supply_rate"] is None and line["install_rate"] is None
                               and line["supply_base_rate"] is None
                               and line["install_base_rate"] is None and rate_mapped)
                    if rate_only and not rate_needed and no_rate:
                        # A rate-only line with no rate on either track (1 October
                        # 2026): the ordinary blocking "rate" need, worded for what
                        # it is, so "Not priced" answers it. Asked whatever the
                        # amount cells hold — its 0 is no quantity, so a 0 amount
                        # does not price it at nil.
                        counts["rate_only_no_rate"] += 1
                        cause = [m for m in line["flags"] if m.startswith(("Supply", "Installation"))]
                        if cause:
                            need("rate", f"{cause[0]} — type a rate")
                        else:
                            msg = flag(rnum, "rate", "ro_no_rate", severity="red")["message"]
                            line["flags"].append(msg)
                            need("rate", msg)
                    # A quantity and no rate anywhere. An EXPLICIT zero amount is
                    # the sheet pricing the line at nil — the Sify schedule's four
                    # nil-priced lines — and is not asked about.
                    elif (not rate_only and not rate_needed and q is not None and no_rate
                            and not amounts):
                        cause = [m for m in line["flags"] if m.startswith(("Supply", "Installation"))]
                        if cause:
                            # The rate cell was already flagged ("#REF!", "NA"): the
                            # need rides on that flag rather than adding a second.
                            need("rate", f"{cause[0]} — type a rate")
                        else:
                            msg = flag(rnum, "rate", "no_rate", severity="red")["message"]
                            line["flags"].append(msg)
                            need("rate", msg)

                if not line["description"]:
                    msg = flag(rnum, "description", "no_desc", severity="red")["message"]
                    line["flags"].append(msg)
                    need("description", msg)
                if not line["item_no"]:
                    msg = flag(rnum, "item_no", "no_item", severity="red")["message"]
                    line["flags"].append(msg)
                    need("item_no", msg)


            # The discount and net-rate cells' own notes (R3) — raised AFTER the
            # rate needs, so a rate need never rides on a discount's flag.
            for f in disc_bad:
                line["flags"].append(f"{TARGET_LABEL[f]}: "
                                     + flag(rnum, f, "disc_range", raw=raws[f])["message"])
            for f in DISC_FIELDS + NET_FIELDS:
                if nflags[f] and nflags[f] != "as_remark":
                    line["flags"].append(f"{TARGET_LABEL[f]}: "
                                         + flag(rnum, f, nflags[f], raw=raws[f])["message"])
            counts["lines"] += 1

            # What this line puts on the running sums — the sheet's own amount
            # where it gave one, quantity × (net) rate where it did not.
            contrib = {}
            for col in amount_cols:
                if col in amounts:
                    contrib[col] = amounts[col]
                    continue
                t = track_of(col)
                qq = line["qty"] or 0.0
                s, i = net_of("supply") or 0.0, net_of("install") or 0.0
                contrib[col] = qq * (s if t == "supply" else i if t == "install" else s + i)
            foot.add(contrib)

        if line["item_src"] == "auto":
            counts["auto_items"] += 1
        elif line["item_src"] == "sheet" and line["kind"] == "sub_item":
            counts["sheet_items"] += 1
        counts["needs"] += len(line["needs"])
        if line["flags"]:
            counts["flagged"] += 1
        lines.append(line)

    # ── Group labels inside an item (1 October 2026) ───────────────────────
    n_labels = _group_labels(lines)
    counts["group_labels"] = n_labels
    counts["headers"] -= n_labels                    # they were counted as spec text

    # ── Section codes ──────────────────────────────────────────────────────
    # The sheet's own codes are reserved first; lines above the first section
    # row take the first free letter (so "A" when the sheet never uses it);
    # a caps-title section with no code takes the next free one after that.
    explicit = [s["code"] for s in sections if s["code"]]
    used = set()
    letters = [chr(c) for c in range(ord("A"), ord("Z") + 1)]

    def next_free():
        for l in letters:
            if l not in used and l not in explicit:
                return l
        n = 1
        while f"S{n}" in used:
            n += 1
        return f"S{n}"

    default_code = ""
    if any(l["section"] == DEFAULT for l in lines):
        default_code = next_free()
        used.add(default_code)

    codes = []
    for s in sections:
        code = s["code"] or next_free()
        if code in used:
            base, n = code, 2
            while f"{base}-{n}" in used:
                n += 1
            new = f"{base}-{n}"
            flag(s["row"], "item_no", "dup_section", raw=base, new=new)
            code = new
        used.add(code)
        codes.append(code)
    out_sections = []
    # Where each section's code came from, aligned with `sections` (6 October
    # 2026, R5): "sheet" when the sheet wrote it, "auto" when the reader
    # assigned it. Several tabs into one BOQ re-letter only the "auto" ones.
    section_src = []
    if default_code:
        out_sections.append({"code": default_code, "title": ""})
        section_src.append("auto")
    for s, code in zip(sections, codes):
        out_sections.append({"code": code, "title": s["title"]})
        section_src.append("sheet" if s["code"] else "auto")
    for l in lines:
        l["section"] = default_code if l["section"] == DEFAULT else codes[l["section"]]
    if not out_sections:
        out_sections.append({"code": "A", "title": ""})
        section_src.append("auto")

    return {"sections": out_sections, "lines": lines, "flags": flags,
            "needs": needs, "checks": checks, "counts": counts,
            "section_src": section_src,
            "totals": _totals_check(lines, totals_rows, col_of)}


def _totals_check(lines: list, totals_rows: list, col_of: dict) -> dict:
    """
    sum(qty × rate) per track against the sheet's own grand total.

    The grand total is the last row labelled "Grand Total" (or "TOTAL (A+B+C)");
    failing that, the one plain "Total" row when there is exactly one. Several
    plain totals and no grand one is reported as no total found — picking one
    would be a guess about which section's figure it is.
    """
    grand = None
    for t in totals_rows:
        if _GRAND.search(t["label"]):
            grand = t
    if grand is None:
        plain = [t for t in totals_rows if not _PARTIAL.search(t["label"])]
        if len(plain) == 1:
            grand = plain[0]

    computed = {"supply": 0.0, "install": 0.0}
    for l in lines:
        if l["is_header"] or l["qty"] is None:
            continue
        for t in ("supply", "install"):
            r = l.get(f"{t}_rate")
            if r is None:
                continue
            # At the NET rate when the line carries a discount (6 October
            # 2026, R3) — what the sheet's own total is the sum of.
            d = l.get(f"{t}_disc_pct")
            computed[t] += l["qty"] * (r * (100.0 - d) / 100.0 if d else r)

    tracks = [t for t in ("supply", "install") if f"{t}_rate" in col_of]
    checks = []
    for t in tracks:
        col = f"{t}_amount" if f"{t}_amount" in col_of else (
            "amount" if "amount" in col_of and len(tracks) == 1 else "")
        checks.append({"track": t, "computed": round(computed[t], 2), "column": col})
    if len(tracks) == 2 and "amount" in col_of and not any(c["column"] == "amount" for c in checks):
        checks.append({"track": "both", "column": "amount",
                       "computed": round(computed["supply"] + computed["install"], 2)})

    status = "no_total" if grand is None else "no_figure"
    for c in checks:
        sheet = grand["amounts"].get(c["column"]) if (grand and c["column"]) else None
        c["sheet"] = None if sheet is None else round(float(sheet), 2)
        if c["sheet"] is None:
            c["status"] = "no_figure"
            continue
        c["difference"] = round(c["computed"] - c["sheet"], 2)
        c["status"] = "match" if abs(c["difference"]) <= TOTALS_TOLERANCE else "mismatch"
    if grand is not None:
        judged = [c for c in checks if c["status"] in ("match", "mismatch")]
        if judged:
            status = "mismatch" if any(c["status"] == "mismatch" for c in judged) else "match"
    return {"status": status, "row": grand["row"] if grand else None,
            "label": grand["label"] if grand else "", "checks": checks,
            "total_rows": len(totals_rows)}


# =============================================================================
# 7. THE ADVICE — what to take from each column (6 October 2026)
# =============================================================================
#
# CLIENT_CHANGES.md §0, forty-fourth block, R1. The client's words: there is
# NO standard BOQ format, and more detection rules will never finish. So the
# preview lists every column with three sample values, an Import tick, a
# target and ONE LINE OF ADVICE, and the tick and the target arrive set to the
# advice. The user decides; this only advises.
#
# `advise()` is PURE and deterministic — a header, its samples and a few facts
# about the rest of the sheet in, `(target, reason)` out — so every rule is a
# unit test and nothing about it depends on a request. `advise_mapping()` works
# the facts out for a whole sheet and calls it once per column. ⚠ It BUILDS ON
# `guess_mapping()`, which is unchanged: the reader's guess is one of the
# facts, so every column the guess already places is advised to stay there.
# What the advice adds is the columns the guess could not judge: a lone rate,
# an empty escalation, a discount, a net rate, a remark, a running total.
#
# ⚠ **A lone rate column is advised as the SELLING rate** — the 29 September
#   2026 decision, kept by the forty-fourth block. `guess_mapping()` still
#   leaves it "?" (`tests/test_boq_import.py` holds that); the advice is what
#   pre-sets it, on a screen the user confirms. A work order has no selling
#   rate, so there a lone rate stays "?" to be chosen (`context["doc"]`).

_RUNNING = re.compile(r"cumulat|running\s*total|progressive|to\s*date|\bupto\b|"
                      r"\bprevious\b|\bc\s*/\s*f\b|\bb\s*/\s*f\b|carried|brought", re.I)
_DISC_HEAD = re.compile(r"\bdisc\b|\bdisc\.|discount|\bless\b|rebate", re.I)
_DISC_AMOUNT = re.compile(r"\bamt\b|amount|\brs\b|\brs\.|₹|\binr\b|\bvalue\b", re.I)
_NET_HEAD = re.compile(r"\bnett?\b|after\s*disc", re.I)
_REMARK_HEAD = re.compile(r"\bremarks?\b|\bnotes?\b|\bcomments?\b|\bobservations?\b", re.I)
_TAX_HEAD = re.compile(r"\bhsn\b|\bsac\b|\bc?gst\b|\bsgst\b|\bigst\b|\btax\b", re.I)

# The reasons, said once. Plain English, no column text in them — a header is
# the sheet's own words and is shown beside the advice, escaped, never inside it.
ADVICE = {
    "empty":     "Nothing in this column. Skip it.",
    "running":   "Running total, not a line value. Skip it.",
    "disc":      "Looks like a discount %. Take it as Discount.",
    "disc_amt":  "Looks like a discount amount, not a %. Map it to Remark.",
    "net":       "The rate after the discount — used only to check rate less discount. Take it.",
    "net_rate":  "Looks like your selling rate (after any discount). Take it.",
    "net_skip":  ("A net rate, but the sheet has no discount column — the unit rate "
                  "already carries the price. Skip it."),
    "remark":    "Notes for the line. Keep them in the remark (never printed).",
    "tax":       "HSN/SAC and GST are not read from the sheet. Skip it.",
    "item_no":   "Item numbers. Take them.",
    "description": "The line text — taken word for word. Take it.",
    "qty":       "The total quantity. Take it.",
    "qty_split": "A floor or area quantity — only the total quantity is read. Skip it.",
    "unit":      "The unit of measure. Take it.",
    "make":      "Brand or make — kept in the line's remark. Take it.",
    "esc":       "Escalation % on the base rate. Take it; it never prints.",
    "esc_none":  "No escalation in this sheet. Skip it; the unit rate stands on its own.",
    "base":      "Base rate — your cost basis. Take it; it never prints.",
    "base_pos":  ("Sits just before the unit rate it is escalated from — looks like the "
                  "base rate. Take it; it never prints."),
    "esc_pos":   ("Sits between a base rate and its unit rate, holding percentages — looks "
                  "like the escalation %. Take it; it never prints."),
    "base_none": "No base rates in this column. Skip it; the unit rate stands on its own.",
    "rate":      "Looks like your selling rate. Take it.",
    "rate_inst": "Looks like your installation rate. Take it.",
    "rate_none": "No rates in this column. Skip it, and price the lines on the form.",
    "lone":      ("Looks like your selling rate. Take it. The sheet does not say Supply "
                  "or Installation — change it if this is installation."),
    "choose":    "A rate, but the sheet does not say Supply or Installation. Choose one.",
    "amount":    "Line amount — used only to check quantity × rate, never imported. Take it.",
    "amount_none": "No amounts in this column. Skip it.",
    "other":     "Not a column the BOQ uses. Skip it.",
    "dup":       "The same target as a column to its left. Skip this one.",
}


def _is_blank(v) -> bool:
    """A data cell that says nothing: empty, whitespace, a dash, or a zero."""
    if v is None:
        return True
    if _is_number(v):
        return abs(v) < 0.005
    s = str(v).strip()
    return not s or bool(_DASH.match(s))


def _as_number(v):
    """A sample as a number — a figure, or text that is one (with an optional
    "%" or thousands commas) — else None."""
    if _is_number(v):
        return float(v)
    s = str(v or "").strip()
    if s.endswith("%"):
        s = s[:-1].strip()
    if _NUMERIC_TEXT.match(s):
        return float(s.replace(",", ""))
    return None


def column_values(grid: dict, ci: int) -> list:
    """Every data cell of grid column `ci`, in sheet order (header excluded)."""
    return [_value(grid, ri, ci) for ri in range(data_start(grid), len(grid["rows"]))]


def column_samples(grid: dict, ci: int, n: int = 3) -> list:
    """The first `n` non-blank data cells of column `ci`, as display text —
    what the preview shows beside each column, as the reader holds the cell."""
    out = []
    for v in column_values(grid, ci):
        if _is_blank(v):
            continue
        if _is_number(v):
            f = float(v)
            out.append(str(int(f)) if f.is_integer() and abs(f) < 1e15 else repr(f))
        else:
            out.append(re.sub(r"\s+", " ", str(v)).strip())
        if len(out) >= n:
            break
    return out


def _track_word(label: str) -> str:
    h = _hits(label)
    return "supply" if "supply" in h else "install" if "installation" in h else ""


def advise(header_text, sample_values, context=None) -> tuple:
    """
    `(suggested_target, reason)` for one column — PURE and deterministic.

    `header_text` is the column's heading (combined for a two-row band);
    `sample_values` its first data cells; `context` the facts about the rest
    of the sheet that `advise_mapping()` works out, every key optional:

      guess        guess_mapping()'s target for this column ("", "?", a key)
      track        the track the heading names: "supply" | "install" | ""
      left_track   the track of the nearest rate column to the LEFT
      tracks       the tracks a unit or base rate is guessed on elsewhere
      all_blank    True when every data cell is blank, a dash or zero
      lone_rate    True for the sheet's only rate-like column, no track named
      net_has_disc True when a discount column sits on this net rate's track
      net_has_rate True when another unit-rate column sits on its track
      doc          "boq" (default) | "wo" — a work order has no selling rate

    The target is a `TARGET_KEYS` member, or `UNDECIDED` where the user must
    choose. The reason is one of `ADVICE`'s sentences.
    """
    ctx = dict(context or {})
    head = str(header_text or "").strip()
    norm = _norm(head)
    samples = [s for s in (sample_values or []) if not _is_blank(s)]
    guess = ctx.get("guess") or ""
    doc = ctx.get("doc") or "boq"
    hits = _hits(head)

    def track_for():
        t = ctx.get("track") or _track_word(head) or ctx.get("left_track") or ""
        if not t:
            tracks = [x for x in (ctx.get("tracks") or []) if x]
            t = tracks[0] if len(set(tracks)) == 1 else "supply"
        return t

    if not head and not samples:
        return "", ADVICE["empty"]
    if guess not in ("item_no", "description") and head and _RUNNING.search(norm):
        return "", ADVICE["running"]

    # A discount — a % per track, or (in rupees) not one at all.
    if guess not in ("item_no", "description", "qty") and head and _DISC_HEAD.search(norm):
        big = any((_as_number(s) or 0.0) > 100 for s in samples)
        if _DISC_AMOUNT.search(norm) or big:
            return "remark", ADVICE["disc_amt"]
        return f"{track_for()}_disc_pct", ADVICE["disc"]

    # A net rate — a check beside a discount, or the selling rate on its own.
    if (guess not in ("item_no", "description", "qty") and head and _NET_HEAD.search(norm)
            and ("rate" in hits or "price" in norm) and "amount" not in hits):
        t = track_for()
        if ctx.get("net_has_disc"):
            return f"{t}_net_rate", ADVICE["net"]
        if not ctx.get("net_has_rate"):
            return f"{t}_rate", ADVICE["net_rate"]
        return "", ADVICE["net_skip"]

    if not guess and head and _REMARK_HEAD.search(norm):
        return "remark", ADVICE["remark"]
    if not guess and head and _TAX_HEAD.search(norm):
        return "", ADVICE["tax"]

    blank = bool(ctx.get("all_blank"))
    if guess in ("item_no", "description", "qty", "unit", "make"):
        return guess, ADVICE[guess]
    if guess in ("escalation_pct", "install_escalation_pct"):
        return ("", ADVICE["esc_none"]) if blank else (guess, ADVICE["esc"])
    if guess in ("supply_base_rate", "install_base_rate"):
        return ("", ADVICE["base_none"]) if blank else (guess, ADVICE["base"])
    if guess in ("supply_rate", "install_rate"):
        if blank:
            return "", ADVICE["rate_none"]
        return guess, ADVICE["rate" if guess == "supply_rate" else "rate_inst"]
    if guess == UNDECIDED:
        if blank:
            return "", ADVICE["esc_none" if "esc" in hits else "rate_none"]
        # Where it SITS: untracked rate columns straight to the left of a
        # tracked unit rate are that rate's base, and the escalation between
        # them (the Sify sheet's "Mohali Rates | Rate increased in % | Supply
        # U/ Rate"). A BOQ only — a work order has neither field.
        if doc == "boq" and ctx.get("base_for"):
            t = ctx["base_for"]
            return f"{t}_base_rate", ADVICE["base_pos"]
        if doc == "boq" and ctx.get("esc_for"):
            t = ctx["esc_for"]
            return ("escalation_pct" if t == "supply" else "install_escalation_pct",
                    ADVICE["esc_pos"])
        if "esc" in hits:
            # An escalation column naming no track, and nowhere telling.
            return UNDECIDED, ADVICE["choose"]
        if doc == "boq" and ctx.get("lone_rate"):
            return "supply_rate", ADVICE["lone"]
        return UNDECIDED, ADVICE["choose"]
    if guess in AMOUNT_FIELDS:
        return ("", ADVICE["amount_none"]) if blank else (guess, ADVICE["amount"])
    if guess:
        return guess, ADVICE["other"]
    if "qty" in hits:
        return "", ADVICE["qty_split"]
    return "", ADVICE["other"]


def advise_mapping(grid: dict, doc: str = "boq") -> dict:
    """
    `{str(Excel column): (target, reason)}` — `advise()` for every column of
    the sheet, with the facts about the sheet worked out once.

    A target advised for two columns keeps the FIRST (left-most) and advises
    skipping the rest — confirm refuses a target chosen twice, and the leftmost
    is what `guess_mapping()` already prefers.
    """
    if grid is None:
        return {}
    cols = grid["cols"]
    labels = header_labels(grid)
    guess = guess_mapping(grid)
    g_of = {ci: guess.get(str(c), "") for ci, c in enumerate(cols)}
    rate_targets = {"supply_rate": "supply", "install_rate": "install",
                    "supply_base_rate": "supply", "install_base_rate": "install"}

    def is_disc(ci):
        n = _norm(labels.get(ci, ""))
        return bool(n) and bool(_DISC_HEAD.search(n)) and g_of[ci] not in (
            "item_no", "description", "qty")

    def is_net(ci):
        lab = labels.get(ci, "")
        n, h = _norm(lab), _hits(lab)
        return (bool(n) and bool(_NET_HEAD.search(n)) and ("rate" in h or "price" in n)
                and "amount" not in h and not is_disc(ci))

    blank = {ci: all(_is_blank(v) for v in column_values(grid, ci)) for ci in range(len(cols))}
    tracks = [rate_targets[g] for ci, g in g_of.items() if g in rate_targets and not blank[ci]]

    def left_track(ci):
        for k in range(ci - 1, -1, -1):
            g = g_of[k]
            if g in ("supply_rate", "install_rate"):
                return rate_targets[g]
            if g == UNDECIDED and not is_net(k) and not blank[k]:
                return "supply" if doc == "boq" else ""
        return ""

    def pct_like(ci) -> bool:
        """Percentages: a cell formatted as one, or every figure 0–100."""
        figs = []
        for ri in range(data_start(grid), len(grid["rows"])):
            v = _value(grid, ri, ci)
            if _kind(grid, ri, ci) == "pct":
                return True
            if _is_number(v) and abs(v) >= 0.005:
                figs.append(float(v))
        return bool(figs) and all(0 <= f <= 100 for f in figs)

    # Untracked rate columns straight to the LEFT of a tracked unit rate, with
    # nothing else between: one is that track's base; two are base, then
    # escalation — when the second holds percentages.
    pos = {}
    if doc == "boq":
        for ci, g in g_of.items():
            if g not in ("supply_rate", "install_rate") or blank[ci]:
                continue
            span, k = [], ci - 1
            while (k >= 0 and g_of[k] == UNDECIDED and not blank[k]
                   and not is_net(k) and not is_disc(k)):
                span.insert(0, k)
                k -= 1
            t = rate_targets[g]
            if len(span) == 1 and "esc" not in _hits(labels.get(span[0], "")):
                pos[span[0]] = ("base_for", t)
            elif len(span) == 2 and pct_like(span[1]):
                pos[span[0]] = ("base_for", t)
                pos[span[1]] = ("esc_for", t)

    undecided = [ci for ci, g in g_of.items()
                 if g == UNDECIDED and not is_net(ci) and not is_disc(ci) and not blank[ci]
                 and "esc" not in _hits(labels.get(ci, "")) and ci not in pos]
    tracked_rate = any(g in ("supply_rate", "install_rate") and not blank[ci]
                       for ci, g in g_of.items())
    lone = len(undecided) == 1 and not tracked_rate

    def track_of_col(ci):
        return _track_word(labels.get(ci, "")) or left_track(ci) or (
            tracks[0] if len(set(tracks)) == 1 else "supply")

    disc_tracks = {track_of_col(ci) for ci in range(len(cols)) if is_disc(ci)}

    out, taken = {}, {}
    for ci, c in enumerate(cols):
        t_here = track_of_col(ci) if (is_disc(ci) or is_net(ci)) else ""
        ctx = {"guess": g_of[ci], "track": _track_word(labels.get(ci, "")),
               "left_track": left_track(ci), "tracks": tracks,
               "all_blank": blank[ci], "lone_rate": lone and ci in undecided,
               "net_has_disc": t_here in disc_tracks,
               "net_has_rate": any(g_of[k] == f"{t_here}_rate" and not blank[k]
                                   for k in range(len(cols)) if k != ci),
               "doc": doc}
        if ci in pos:
            ctx[pos[ci][0]] = pos[ci][1]
        target, reason = advise(labels.get(ci, ""), column_samples(grid, ci), ctx)
        if target and target != UNDECIDED and target in taken:
            target, reason = "", ADVICE["dup"]
        if target and target != UNDECIDED:
            taken[target] = ci
        out[str(c)] = (target, reason)
    return out


def advised_mapping(grid: dict, doc: str = "boq") -> dict:
    """`{str(column): target}` — the advice's targets alone: the mapping a
    freshly staged sheet starts from (R1 — the tick and the target are
    pre-set to the advice)."""
    return {c: t for c, (t, _r) in advise_mapping(grid, doc).items()}


# =============================================================================
# 8. NAMES FROM THE SHEET — the project and the customer (6 October 2026, R4)
# =============================================================================
#
# SUGGESTIONS the user confirms on the preview, never saved by themselves.
# Only the rows ABOVE the heading are read: that is where a sheet writes its
# title block. Each name comes with the row it was found on, so the preview
# can say where. Nothing is invented: a sheet with no such row gives None.
#
# ⚠ **"Name of Work" is the PROJECT, not the customer** — the brief listed it
#   among the client labels, but on an Indian tender BOQ it is the name of the
#   work itself ("Name of Work: Fire fighting system at …"); read as the
#   customer it would put the job's description in the bill-to box. It is
#   read as the project title, and the pass report says so.

_PARTY = r"(?:client|customer|owner|party|contractor|employer|buyer)"
_CLIENT_LABEL = re.compile(
    r"^\s*(?:name\s+of\s+(?:the\s+)?)?" + _PARTY +
    r"(?:'s)?\s*(?:name)?\s*(?:[:\-–—]+\s*(?P<v>.*))?$", re.I | re.S)
_MS_PREFIX = re.compile(r"^\s*m\s*/\s*s\.?\s*(?P<v>\S.*)$", re.I | re.S)
_PROJECT_LABEL = re.compile(
    r"^\s*(?:name\s+of\s+(?:the\s+)?)?(?:project|work)\s*(?:name|title)?"
    r"\s*(?:[:\-–—]+\s*(?P<v>.*))?$", re.I | re.S)
# A document title that names no job — "BILL OF QUANTITIES", "BOQ", "PRICE
# SCHEDULE". Skipped as a project title; "BOQ for Fire Fighting at X" is not.
_GENERIC_TITLE = re.compile(r"\b(?:bill\s+of\s+quantit(?:y|ies)|b\.?\s*o\.?\s*q\.?|"
                            r"schedule\s+of\s+(?:rates|quantities)|price\s+schedule|"
                            r"abstract|annexure|rate\s+schedule)\b|[^a-z]", re.I)


def _texts(vals: list) -> list:
    """`[(column index, text)]` of a row's non-empty text cells, in order."""
    return [(ci, str(v).strip()) for ci, v in enumerate(vals)
            if isinstance(v, str) and v.strip()]


def _labelled(cells: list, rx) -> str:
    """The value of the first cell matching a `Label: value` pattern — after
    the colon, or the next text cell on the row — else ""."""
    for k, (_ci, text) in enumerate(cells):
        m = rx.match(text)
        if not m:
            continue
        value = (m.group("v") or "").strip()
        if not value and k + 1 < len(cells):
            value = cells[k + 1][1]
        if value:
            return value
    return ""


def detect_names(grid: dict) -> dict:
    """
    `{"project": {"text", "row", "how"} | None, "client": {...} | None}` —
    read from the rows above the heading, word for word (trimmed only).

    * **client** — a cell labelled Client / Customer / Owner / Party /
      Contractor / Employer / Buyer (also "Name of …", with or without ":"),
      its value after the colon or in the next cell on the row; else a cell
      that opens "M/s", the whole cell (`how`: "label" / "ms").
    * **project** — a cell labelled Project / Name of Work, its value; else
      the first row holding ONE text cell (a merged title) that is not a bare
      document title like "BILL OF QUANTITIES" (`how`: "label" / "title").

    The tab's name is not a guess here: the caller offers it as the fallback,
    and says that is what it is.
    """
    out = {"project": None, "client": None}
    hdr = (grid or {}).get("header") or []
    if not hdr:
        return out
    title = None
    for ri in range(0, hdr[0]):
        rnum, vals = grid["rows"][ri]
        cells = _texts(vals)
        if not cells:
            continue
        if out["client"] is None:
            v, how = _labelled(cells, _CLIENT_LABEL), "label"
            if not v:
                # "M/s Acme Infra" — the honorific is part of what the sheet
                # wrote, so the cell is taken whole, word for word.
                for _ci, text in cells:
                    if _MS_PREFIX.match(text):
                        v, how = text, "ms"
                        break
            if v:
                out["client"] = {"text": v, "row": rnum, "how": how}
                continue
        if out["project"] is None:
            v = _labelled(cells, _PROJECT_LABEL)
            if v:
                out["project"] = {"text": v, "row": rnum, "how": "label"}
                continue
        if title is None and len(cells) == 1:
            text = cells[0][1]
            if (len(_GENERIC_TITLE.sub("", text)) >= 3 and not _CLIENT_LABEL.match(text)
                    and not _MS_PREFIX.match(text) and not _PROJECT_LABEL.match(text)):
                title = {"text": text, "row": rnum, "how": "title"}
    if out["project"] is None and title is not None:
        out["project"] = title
    return out
